#!/usr/bin/env python3
"""Check OPPO serial numbers/IMEIs on official OPPO Support Ukraine portal.

Uses local OpenCV image processing (slider_cv.py) for slider CAPTCHA solving:
- Canny edge detection & template matching on puzzle pieces
- Fallback morphological contour shadow detection
- Real-time progress bar via tqdm
- Formatted Excel updating matching oppo_results.xlsx
"""

from __future__ import annotations

import argparse
import base64
import datetime
import math
import random
import sys
import os
import csv
import time
from io import BytesIO
from pathlib import Path
from typing import Any

# Ensure Playwright finds shared system browsers (crucial for PyInstaller EXE)
if "PLAYWRIGHT_BROWSERS_PATH" not in os.environ or not os.environ["PLAYWRIGHT_BROWSERS_PATH"]:
    _local_appdata = os.environ.get("LOCALAPPDATA") or str(Path.home() / "AppData" / "Local")
    os.environ["PLAYWRIGHT_BROWSERS_PATH"] = str(Path(_local_appdata) / "ms-playwright")

from playwright.sync_api import sync_playwright

import cv2
from slider_cv import (
    decode_image_bytes,
    detect_hole_by_contours,
    find_slider_displacement,
    match_template_edges,
)


try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    from tqdm import tqdm
    HAS_TQDM = True
except ImportError:
    HAS_TQDM = False

# Ensure real-time console streaming
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(line_buffering=True, encoding="utf-8")
    except Exception:
        pass


def read_identifiers(file_path: Path) -> list[str]:
    """Read serial numbers / IMEIs from a text or excel/csv file."""
    if not file_path.exists():
        raise FileNotFoundError(f"Input file not found: {file_path}")

    ext = file_path.suffix.lower()
    identifiers: list[str] = []

    if ext in [".xlsx", ".xls"] and HAS_PANDAS:
        df = pd.read_excel(file_path)
        first_col = df.columns[0]
        for val in df[first_col].dropna():
            s = str(val).strip()
            if s and not s.startswith("#"):
                identifiers.append(s)
        return identifiers

    for line in file_path.read_text(encoding="utf-8-sig").splitlines():
        val = line.strip()
        if val and not val.startswith("#"):
            identifiers.append(val)
    return identifiers


def format_oppo_date(raw_val: Any) -> str:
    """Format timestamp into readable date matching OPPO portal, e.g. '25-09-2026 (UTC+3)'."""
    if not raw_val:
        return ""
    if isinstance(raw_val, (datetime.datetime, datetime.date)):
        return raw_val.strftime("%d-%m-%Y (UTC+3)")
    s = str(raw_val).strip()
    if not s or s in ["Not Activated", "Pending Sync", "Unknown", "Not Found"]:
        return s
    # Already formatted like DD-MM-YYYY (UTC+3)
    if "(UTC" in s:
        return s

    # Check if numeric epoch timestamp in ms (13 digits) or seconds (10 digits)
    try:
        val_float = float(s)
        if val_float > 1_000_000_000:
            if val_float > 10_000_000_000:  # millisecond epoch
                val_float = val_float / 1000.0
            tz = datetime.timezone(datetime.timedelta(hours=3))
            dt = datetime.datetime.fromtimestamp(val_float, tz)
            return dt.strftime("%d-%m-%Y (UTC+3)")
    except Exception:
        pass

    # If string date formatted as YYYY-MM-DD or DD-MM-YYYY
    if "-" in s and len(s) >= 10:
        try:
            date_part = s.split(" ")[0].strip()
            parts = date_part.split("-")
            if len(parts) == 3 and len(parts[0]) == 4:  # YYYY-MM-DD -> DD-MM-YYYY (UTC+3)
                return f"{parts[2]}-{parts[1]}-{parts[0]} (UTC+3)"
            if len(parts) == 3 and len(parts[0]) == 2 and len(parts[2]) == 4:  # DD-MM-YYYY -> DD-MM-YYYY (UTC+3)
                return f"{parts[0]}-{parts[1]}-{parts[2]} (UTC+3)"
        except Exception:
            pass
    return s


def generate_natural_track(distance: int) -> list[int]:
    """Generate realistic human dragging track with natural velocity curve."""
    steps = random.randint(26, 36)
    track: list[int] = []

    for i in range(1, steps + 1):
        t = i / steps
        eased = 0.5 * (1 - math.cos(math.pi * t))
        track.append(round(distance * eased))

    deltas = [track[0]]
    for i in range(1, len(track)):
        deltas.append(track[i] - track[i - 1])

    return deltas


def solve_dingxiang_captcha_cv(page, max_attempts: int = 7, on_status: Any = None, fine_tune_offset: float = 0.0) -> bool:
    """Automatically solve Dingxiang slider CAPTCHA using dual-strategy closed-loop OpenCV processing."""
    bar = page.locator("[id*='dx_captcha_basic_bar']").first

    # Wait up to 5 seconds for captcha to appear
    for _ in range(25):
        if bar.is_visible():
            break
        page.wait_for_timeout(200)

    if not bar.is_visible():
        return True

    if on_status:
        on_status("Captcha detected, processing image locally...")

    offsets = [0.0, 1.0, -1.0, 2.0, -2.0, 3.0, -3.0]

    for attempt in range(max_attempts):
        if not bar.is_visible():
            if on_status:
                on_status("Captcha solved!")
            return True

        if on_status:
            on_status(f"OpenCV solve attempt {attempt + 1}/{max_attempts}...")

        # Wait for canvas to draw
        page.wait_for_timeout(1500)

        # Extract canvas and piece image sources plus exact geometry ratios from live DOM
        puzzle_data = page.evaluate("""() => {
            const canvas = document.querySelector('#dx_captcha_basic_bg_1 canvas') || document.querySelector('.dx_captcha_basic_bg canvas');
            const subImg = document.querySelector('[id*="dx_captcha_basic_sub-slider"] img') || document.querySelector('.dx_captcha_basic_sub-slider img');
            const sub = document.querySelector('[id*="dx_captcha_basic_sub-slider"]') || document.querySelector('.dx_captcha_basic_sub-slider');
            const slider = document.querySelector('[id*="dx_captcha_basic_slider_"]');
            const cRect = canvas ? canvas.getBoundingClientRect() : null;
            const sRect = slider ? slider.getBoundingClientRect() : null;
            return {
                bg: canvas ? canvas.toDataURL('image/png') : null,
                pieceSrc: subImg ? subImg.src : null,
                marginTop: sub ? sub.style.marginTop : '0px',
                scaleX: (canvas && cRect) ? (cRect.width / canvas.width) : 0.72,
                sliderBox: sRect ? {x: sRect.left, y: sRect.top, width: sRect.width, height: sRect.height} : null
            };
        }""")

        canvas_data = puzzle_data.get("bg")
        if not canvas_data:
            page.wait_for_timeout(1000)
            continue

        try:
            raw_bg_bytes = base64.b64decode(canvas_data.split(",")[1])
            piece_src = puzzle_data.get("pieceSrc")
            raw_piece_bytes = None

            if piece_src:
                try:
                    resp = page.request.get(piece_src)
                    if resp.status == 200:
                        raw_piece_bytes = resp.body()
                except Exception:
                    raw_piece_bytes = None

            margin_top_str = puzzle_data.get("marginTop", "0px")
            margin_top = float(str(margin_top_str).replace("px", "").strip() or 0)
            dynamic_scale = float(puzzle_data.get("scaleX") or 0.72)
            target_y_400 = int(round(margin_top / dynamic_scale)) if margin_top else None

            bg_cv = decode_image_bytes(raw_bg_bytes, cv2.IMREAD_COLOR)

            # Dual-strategy: Template match (LAB) + Contour detection (CLAHE)
            xm, ym, cm = (0, 0, 0.0)
            if raw_piece_bytes:
                piece_cv = decode_image_bytes(raw_piece_bytes, cv2.IMREAD_UNCHANGED)
                xm, ym, cm = match_template_edges(bg_cv, piece_cv)

            xc, yc, cc = detect_hole_by_contours(bg_cv, target_y=target_y_400)

            if cm >= 0.25:
                best_x_400 = xm
            elif cc > 50:
                best_x_400 = xc
            else:
                best_x_400 = xm if cm > 0 else xc

            offset = offsets[attempt % len(offsets)] + fine_tune_offset
            target_hole_css = (best_x_400 * dynamic_scale) + offset
        except Exception as err:
            if on_status:
                on_status(f"CV error: {err}")
            page.wait_for_timeout(1000)
            continue


        slider_box = puzzle_data.get("sliderBox")
        if not slider_box:
            bar_box = bar.bounding_box()
            if not bar_box:
                page.wait_for_timeout(1000)
                continue
            start_x = bar_box["x"] + 22
            start_y = bar_box["y"] + 15
        else:
            start_x = slider_box["x"] + slider_box["width"] / 2
            start_y = slider_box["y"] + slider_box["height"] / 2

        page.mouse.move(start_x, start_y)
        time.sleep(random.uniform(0.1, 0.15))
        page.mouse.down()
        time.sleep(random.uniform(0.1, 0.15))

        cur_x = start_x
        step = random.uniform(2.8, 3.8)
        while True:
            cur_x += step
            wobble = random.uniform(-0.5, 0.5)
            page.mouse.move(cur_x, start_y + wobble)
            time.sleep(random.uniform(0.008, 0.015))

            sub_pos = page.evaluate("""() => {
                const sub = document.querySelector('[id*="dx_captcha_basic_sub-slider"]');
                const canvas = document.querySelector('#dx_captcha_basic_bg_1 canvas');
                return (sub && canvas) ? (sub.getBoundingClientRect().left - canvas.getBoundingClientRect().left) : 0;
            }""")

            if sub_pos >= target_hole_css:
                break
            if cur_x - start_x > 245:
                break

        time.sleep(random.uniform(0.15, 0.25))
        page.mouse.up()

        # Wait for verify response
        page.wait_for_timeout(3500)

        if not bar.is_visible():
            if on_status:
                on_status("Captcha verified!")
            return True

        # Click refresh button for next attempt
        refresh_btn = page.locator("[id*='btn-refresh']").first
        if refresh_btn.is_visible():
            try:
                refresh_btn.click()
            except Exception:
                pass
            for _ in range(25):
                page.wait_for_timeout(200)
                new_data = page.evaluate("""() => {
                    const canvas = document.querySelector('#dx_captcha_basic_bg_1 canvas') || document.querySelector('.dx_captcha_basic_bg canvas');
                    return canvas ? canvas.toDataURL('image/png') : null;
                }""")
                if new_data and new_data != canvas_data:
                    break
        page.wait_for_timeout(1000)

    return not bar.is_visible()


def check_oppo_serial(page, serial: str, on_status: Any = None, fine_tune_offset: float = 0.0) -> dict[str, Any]:
    """Submit a single serial number and extract device verification information."""
    captured_payload: dict[str, Any] = {}

    def on_response(response):
        if "getdeviceinfo" in response.url.lower():
            try:
                captured_payload["data"] = response.json()
            except Exception:
                pass

    page.on("response", on_response)

    try:
        if on_status:
            on_status("Loading portal...")
        page.goto("https://support.oppo.com/ua/check/", wait_until="domcontentloaded", timeout=40000)
        page.wait_for_timeout(2500)

        input_box = page.locator("input.el-input__inner").first
        input_box.fill(serial)
        page.wait_for_timeout(400)

        if on_status:
            on_status("Submitting serial...")
        search_btn = page.locator(".searchBtn").first
        search_btn.click()
        page.wait_for_timeout(2000)

        captcha_solved = solve_dingxiang_captcha_cv(page, on_status=on_status, fine_tune_offset=fine_tune_offset)

        if not captcha_solved:
            return {
                "serial": serial,
                "is_active": "Unknown",
                "activation_date": "",
                "activation_comment": "",
                "model_name": "",
                "specs": "",
                "color": "",
                "certified_model": "",
                "region": "",
                "warranty_end": "",
                "status": "Captcha Failed",
                "raw_msg": "Captcha could not be resolved automatically",
            }

        # Wait for API or DOM to update
        page.wait_for_timeout(4000)

        # Extract live DOM text from the result card for activation time
        dom_val = ""
        try:
            dom_val = page.evaluate("""() => {
                const items = Array.from(document.querySelectorAll('li, div.flex, .result-item'));
                for (let item of items) {
                    const label = item.querySelector('.label');
                    const val = item.querySelector('.value');
                    if (label && val && (label.innerText || '').toLowerCase().includes('активац')) {
                        return (val.innerText || '').trim();
                    }
                }
                const all = Array.from(document.querySelectorAll('*'));
                for (let el of all) {
                    const txt = (el.innerText || '');
                    if (txt.includes('Час активації：') || txt.includes('Час активації:')) {
                        const parts = txt.split(/Час активації[：:]/);
                        if (parts.length > 1) {
                            return parts[1].split('\\n')[0].trim();
                        }
                    }
                }
                return '';
            }""") or ""
        except Exception:
            dom_val = ""

        # Check API response first
        api_data = captured_payload.get("data")
        if isinstance(api_data, dict):
            code = str(api_data.get("code", ""))
            d = api_data.get("data") or {}

            if code == "1" and d:
                reg_date = d.get("regDate")
                pmsg = str(d.get("pmsg") or "")

                dom_lower = dom_val.lower()
                pmsg_lower = pmsg.lower()

                # Determine is_active, activation_date, and activation_comment
                if (
                    "передано до системи" in dom_lower
                    or "7 днів" in dom_lower
                    or "7 days" in dom_lower
                    or "transferred" in dom_lower
                    or "上报系统" in pmsg_lower
                    or "7天" in pmsg_lower
                ):
                    is_active = "Pending"
                    act_date = "Pending Sync"
                    act_comment = "The activation time has been transferred to the system. Please check again after 7 days."
                elif (
                    "не активовано" in dom_lower
                    or "not activat" in dom_lower
                    or "注册时间为空" in pmsg_lower
                    or (not reg_date and not dom_val)
                ):
                    is_active = "No"
                    act_date = "Not Activated"
                    act_comment = "Not Activated"
                elif dom_val and "передано" not in dom_lower:
                    is_active = "Yes"
                    act_date = format_oppo_date(dom_val)
                    act_comment = act_date
                elif reg_date:
                    is_active = "Yes"
                    act_date = format_oppo_date(reg_date)
                    act_comment = act_date
                elif dom_val:
                    is_active = "Yes"
                    act_date = format_oppo_date(dom_val)
                    act_comment = act_date
                else:
                    is_active = "No"
                    act_date = "Not Activated"
                    act_comment = pmsg or "Not Activated"

                ram = str(d.get("productRam") or "").strip()
                rom = str(d.get("productRom") or "").strip()
                specs = f"{ram}+{rom}" if ram and rom else (ram or rom)

                return {
                    "serial": serial,
                    "is_active": is_active,
                    "activation_date": act_date,
                    "activation_comment": act_comment,
                    "model_name": str(d.get("marketingModelName") or "").strip(),
                    "specs": specs,
                    "color": str(d.get("productColor") or "").strip(),
                    "certified_model": str(d.get("certifiedModel") or "").strip(),
                    "region": str(d.get("region") or "").strip(),
                    "warranty_end": format_oppo_date(str(d.get("warrantyEndTime") or "").strip()),
                    "status": "Found",
                    "raw_msg": pmsg,
                }
            elif code != "1":
                msg = str(api_data.get("msg") or "Device not found")
                return {
                    "serial": serial,
                    "is_active": "Not Found",
                    "activation_date": "",
                    "activation_comment": "",
                    "model_name": "",
                    "specs": "",
                    "color": "",
                    "certified_model": "",
                    "region": "",
                    "warranty_end": "",
                    "status": "Not Found",
                    "raw_msg": msg,
                }

        # Fallback to DOM inspection
        text = page.inner_text("body")
        if "Результат пошуку" in text:
            dom_lower = dom_val.lower()
            if "передано до системи" in dom_lower or "7 днів" in dom_lower or "7 days" in dom_lower:
                is_active = "Pending"
                act_date = "Pending Sync"
                act_comment = "The activation time has been transferred to the system. Please check again after 7 days."
            elif "не активовано" in dom_lower or "не активовано" in text.lower():
                is_active = "No"
                act_date = "Not Activated"
                act_comment = "Not Activated"
            elif dom_val:
                is_active = "Yes"
                act_date = format_oppo_date(dom_val)
                act_comment = act_date
            else:
                is_active = "Yes"
                act_date = "Activated"
                act_comment = "Activated"

            return {
                "serial": serial,
                "is_active": is_active,
                "activation_date": act_date,
                "activation_comment": act_comment,
                "model_name": "",
                "specs": "",
                "color": "",
                "certified_model": "",
                "region": "",
                "warranty_end": "",
                "status": "Found (DOM fallback)",
                "raw_msg": "Extracted from page text",
            }
        elif "не знайдено" in text.lower():
            return {
                "serial": serial,
                "is_active": "Not Found",
                "activation_date": "",
                "activation_comment": "",
                "model_name": "",
                "specs": "",
                "color": "",
                "certified_model": "",
                "region": "",
                "warranty_end": "",
                "status": "Not Found",
                "raw_msg": "Device not found in system",
            }

        return {
            "serial": serial,
            "is_active": "Unknown",
            "activation_date": "",
            "activation_comment": "",
            "model_name": "",
            "specs": "",
            "color": "",
            "certified_model": "",
            "region": "",
            "warranty_end": "",
            "status": "No Response",
            "raw_msg": "Neither API response nor result card found",
        }
    except Exception as exc:
        return {
            "serial": serial,
            "is_active": "Error",
            "activation_date": "",
            "activation_comment": "",
            "model_name": "",
            "specs": "",
            "color": "",
            "certified_model": "",
            "region": "",
            "warranty_end": "",
            "status": "Error",
            "raw_msg": str(exc),
        }
    finally:
        try:
            page.remove_listener("response", on_response)
        except Exception:
            pass


def save_to_excel(results: list[dict[str, Any]], output_path: Path, quiet: bool = False, fresh: bool = False) -> dict[str, Any]:
    """Save or update results in Excel spreadsheet and companion CSV with clean formatting."""
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    existing_rows: dict[str, dict[str, Any]] = {}

    columns = [
        ("serial", "Serial / IMEI"),
        ("is_active", "Active?"),
        ("activation_date", "Activation Date"),
        ("activation_comment", "Activation Comment"),
        ("model_name", "Model Name"),
        ("specs", "RAM/ROM"),
        ("color", "Color"),
        ("certified_model", "Certified Model"),
        ("region", "Region"),
        ("warranty_end", "Warranty End"),
        ("status", "Check Status"),
        ("checked_at", "Last Checked"),
    ]

    # Only load existing rows if NOT in fresh mode
    if not fresh and output_path.exists() and HAS_OPENPYXL:
        try:
            wb_old = openpyxl.load_workbook(output_path)
            ws_old = wb_old.active
            headers = [cell.value for cell in ws_old[1]]
            header_map = {col_label: col_key for col_key, col_label in columns}

            for row in ws_old.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                row_dict = {}
                for idx, cell_val in enumerate(row):
                    if idx < len(headers) and headers[idx] in header_map:
                        row_dict[header_map[headers[idx]]] = cell_val
                sn = str(row[0]).strip()
                if sn:
                    # Normalize date formats if older rows had raw millisecond epoch timestamps
                    if row_dict.get("activation_date"):
                        row_dict["activation_date"] = format_oppo_date(row_dict["activation_date"])
                    if row_dict.get("activation_comment"):
                        row_dict["activation_comment"] = format_oppo_date(row_dict["activation_comment"])
                    if row_dict.get("warranty_end"):
                        row_dict["warranty_end"] = format_oppo_date(row_dict["warranty_end"])

                    # Backfill activation_comment for older rows if missing
                    if not row_dict.get("activation_comment"):
                        act_d = str(row_dict.get("activation_date") or "").strip()
                        if act_d and act_d != "Not Activated":
                            row_dict["activation_comment"] = act_d
                        elif row_dict.get("is_active") == "No":
                            row_dict["activation_comment"] = "Not Activated"
                    existing_rows[sn] = row_dict
        except Exception as e:
            if not quiet:
                print(f"[Notice] Could not read existing Excel file: {e}")

    # Merge or update rows by serial number (in-place edit)
    for r in results:
        sn = str(r["serial"]).strip()
        new_status = str(r.get("status") or "")
        old_status = str(existing_rows.get(sn, {}).get("status") or "")
        if sn in existing_rows and old_status == "Found" and new_status != "Found":
            continue
        row_entry = {**r, "checked_at": timestamp}
        existing_rows[sn] = row_entry

    final_rows = list(existing_rows.values())

    # 1. Save to Excel (.xlsx)
    xlsx_path = output_path.with_suffix(".xlsx") if output_path.suffix.lower() == ".csv" else output_path
    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "OPPO Status"

        # Headers
        header_fill = PatternFill(start_color="107C41", end_color="107C41", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_side = Side(border_style="thin", color="D9D9D9")
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[1].height = 28

        for col_idx, (_, col_title) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Rows
        for row_idx, data_row in enumerate(final_rows, start=2):
            ws.row_dimensions[row_idx].height = 22
            is_active_val = data_row.get("is_active", "")

            status_fill = None
            if is_active_val == "Yes":
                status_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
            elif is_active_val == "No":
                status_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            elif is_active_val in ["Pending", "Syncing"]:
                status_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
            elif is_active_val in ["Not Found", "Error", "Captcha Failed"]:
                status_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

            for col_idx, (col_key, _) in enumerate(columns, start=1):
                val = data_row.get(col_key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = cell_border

                if col_key in ["serial", "model_name", "color", "activation_comment"]:
                    cell.alignment = left_align
                else:
                    cell.alignment = center_align

                if col_key == "is_active" and status_fill:
                    cell.fill = status_fill
                    cell.font = Font(name="Calibri", size=10, bold=True)

        for col_idx in range(1, len(columns) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, len(final_rows) + 2))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        wb.save(xlsx_path)
        if not quiet:
            print(f"Successfully saved {len(final_rows)} records to Excel: {xlsx_path}")
    elif HAS_PANDAS:
        df = pd.DataFrame(final_rows)
        df.to_excel(xlsx_path, index=False)
        if not quiet:
            print(f"Saved {len(final_rows)} records to Excel: {xlsx_path}")

    # 2. Save companion CSV (.csv) simultaneously
    csv_path = output_path.with_suffix(".csv")
    try:
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f_csv:
            writer = csv.DictWriter(f_csv, fieldnames=[col_title for _, col_title in columns])
            writer.writeheader()
            for r_entry in final_rows:
                writer.writerow({col_title: r_entry.get(col_key, "") for col_key, col_title in columns})
        if not quiet:
            print(f"Saved companion CSV: {csv_path}")
    except Exception as e_csv:
        if not quiet:
            print(f"[Notice] Could not save CSV file: {e_csv}")

    # Calculate exact summary matching the sheet 1:1
    summary = {
        "total": len(final_rows),
        "yes": sum(1 for r in final_rows if str(r.get("is_active") or "").strip() == "Yes"),
        "no": sum(1 for r in final_rows if str(r.get("is_active") or "").strip() == "No"),
        "pending": sum(1 for r in final_rows if str(r.get("is_active") or "").strip() in ["Pending", "Syncing"]),
        "failed": sum(
            1 for r in final_rows
            if str(r.get("status") or "").strip() in ["Captcha Failed", "Error", "Not Found"]
            or str(r.get("is_active") or "").strip() in ["Not Found", "Error", "Captcha Failed"]
        ),
        "rows": final_rows,
    }
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="Check OPPO serial/IMEI using local OpenCV slider processing.")
    parser.add_argument("--input", type=Path, default=Path("serials.txt"), help="Path to input serials file (default: serials.txt)")
    parser.add_argument("--output", type=Path, default=Path("oppo_results.xlsx"), help="Path to output Excel file (default: oppo_results.xlsx)")
    parser.add_argument("--headful", action="store_true", help="Launch visible browser window instead of headless mode")
    parser.add_argument("--delay", type=float, default=2.0, help="Delay between checks in seconds (default: 2.0)")
    parser.add_argument("--skip-found", action="store_true", help="Skip serials that are already successfully verified in output Excel")
    parser.add_argument("--limit", type=int, default=None, help="Limit number of serials to process (optional)")
    parser.add_argument("--adjust", type=float, default=0.0, help="Nudge drag distance in pixels (+ moves right, - moves left, default: 0.0)")
    args = parser.parse_args()

    identifiers = read_identifiers(args.input)

    if args.skip_found and args.output.exists() and HAS_OPENPYXL:
        try:
            wb_check = openpyxl.load_workbook(args.output)
            ws_check = wb_check.active
            headers = [cell.value for cell in ws_check[1]]
            status_idx = headers.index("Check Status") if "Check Status" in headers else -1
            sn_idx = headers.index("Serial / IMEI") if "Serial / IMEI" in headers else 0
            found_sns = set()
            for row in ws_check.iter_rows(min_row=2, values_only=True):
                if row and status_idx != -1 and len(row) > status_idx and row[status_idx] == "Found":
                    found_sns.add(str(row[sn_idx]).strip())
            initial_count = len(identifiers)
            identifiers = [s for s in identifiers if s not in found_sns]
            if initial_count != len(identifiers):
                print(f"Skipping {initial_count - len(identifiers)} already verified serial(s). {len(identifiers)} remaining to check.")
        except Exception:
            pass

    if args.limit:
        identifiers = identifiers[: args.limit]


    if not identifiers:
        print(f"All serials in {args.input} are already verified in {args.output}!")
        return 0

    print(f"Loaded {len(identifiers)} serial(s) to check via OpenCV Slider Solver.")
    print(f"Target Excel output: {args.output}\n")

    results: list[dict[str, Any]] = []

    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(headless=not args.headful)
        except Exception:
            try:
                browser = p.chromium.launch(channel="msedge", headless=not args.headful)
            except Exception:
                browser = p.chromium.launch(channel="chrome", headless=not args.headful)
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            viewport={"width": 1280, "height": 800},
        )
        page = context.new_page()

        if HAS_TQDM:
            pbar = tqdm(
                identifiers,
                desc="Checking OPPO (OpenCV)",
                unit="device",
                ncols=95,
                bar_format="{l_bar}{bar}| {n_fmt}/{total_fmt} [{elapsed}<{remaining}] {postfix}",
            )
            for idx, sn in enumerate(pbar, start=1):
                pbar.set_postfix_str(f"{sn}: Connecting...")
                res = check_oppo_serial(page, sn, on_status=lambda msg: pbar.set_postfix_str(f"{sn}: {msg}"), fine_tune_offset=args.adjust)
                results.append(res)

                save_to_excel(results, args.output, quiet=True)

                status_line = f"[{idx}/{len(identifiers)}] {sn} -> Active: {res['is_active']} | Model: {res['model_name']} | Color: {res['color']}"
                tqdm.write(status_line)
                pbar.set_postfix_str(f"Last: {res['is_active']}")

                if idx < len(identifiers):
                    time.sleep(args.delay)
            pbar.close()
        else:
            for idx, sn in enumerate(identifiers, start=1):
                pct = int((idx / len(identifiers)) * 100)
                bar_str = "[" + "=" * (pct // 5) + " " * (20 - pct // 5) + "]"
                print(f"{bar_str} {pct}% ({idx}/{len(identifiers)}) Checking: {sn} ...", flush=True)

                res = check_oppo_serial(page, sn, on_status=lambda msg: print(f"   [{sn}] {msg}", flush=True), fine_tune_offset=args.adjust)
                results.append(res)
                save_to_excel(results, args.output, quiet=True)

                print(f"   Result: Active={res['is_active']} | Model={res['model_name']} | Color={res['color']}", flush=True)
                if idx < len(identifiers):
                    time.sleep(args.delay)

        browser.close()

    save_to_excel(results, args.output, quiet=False)
    print("\n" + "=" * 60)
    print(f"Done! Checked {len(results)} device(s). Results written to {args.output}")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

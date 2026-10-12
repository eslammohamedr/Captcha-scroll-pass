#!/usr/bin/env python3
"""Check OPPO serial numbers/IMEIs on official OPPO Support Ukraine portal.

Site: https://support.oppo.com/ua/check/
Extracts device activation status, model, color, specs, and updates an Excel spreadsheet.
"""

from __future__ import annotations

import argparse
import base64
import datetime
import math
import random
import sys
import time
from io import BytesIO
from pathlib import Path
from typing import Any

import numpy as np
from PIL import Image
from playwright.sync_api import sync_playwright

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    from tqdm import tqdm
    HAS_TQDM = True
except ImportError:
    HAS_TQDM = False

# Ensure real-time console streaming
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(line_buffering=True, encoding="utf-8")
    except Exception:
        pass


def read_identifiers(file_path: Path) -> list[str]:
    """Read serial numbers / IMEIs from a text or excel/csv file."""
    if not file_path.exists():
        raise FileNotFoundError(f"Input file not found: {file_path}")

    ext = file_path.suffix.lower()
    identifiers: list[str] = []

    if ext in [".xlsx", ".xls"] and HAS_PANDAS:
        df = pd.read_excel(file_path)
        first_col = df.columns[0]
        for val in df[first_col].dropna():
            s = str(val).strip()
            if s and not s.startswith("#"):
                identifiers.append(s)
        return identifiers

    # Default to text file
    for line in file_path.read_text(encoding="utf-8-sig").splitlines():
        val = line.strip()
        if val and not val.startswith("#"):
            identifiers.append(val)
    return identifiers


def find_hole_x_in_288(canvas_bytes: bytes, margin_top_px: float) -> float:
    """Analyze canvas image and locate cutout center X in 288px coordinate space."""
    img = Image.open(BytesIO(canvas_bytes)).convert("RGB")
    width, height = img.size
    arr = np.array(img, dtype=np.float32)

    gray = 0.299 * arr[:, :, 0] + 0.587 * arr[:, :, 1] + 0.114 * arr[:, :, 2]

    scale_y = height / 144.0
    scale_x = width / 288.0

    target_y = int((margin_top_px + 21) * scale_y)
    target_y = max(20, min(height - 20, target_y))

    radius = int(18 * scale_x)
    y_min = max(0, target_y - radius)
    y_max = min(height, target_y + radius)

    best_canvas_x = None
    min_val = float("inf")

    # The hole is always to the right of the starting piece (x > 50px)
    for x in range(int(55 * scale_x), width - radius):
        window = gray[y_min:y_max, x - radius : x + radius]
        val = float(np.mean(window))
        if val < min_val:
            min_val = val
            best_canvas_x = x

    if best_canvas_x is None:
        best_canvas_x = int(width / 2)

    return float(best_canvas_x / scale_x)


def generate_natural_track(distance: int) -> list[int]:
    """Generate realistic human dragging track with natural velocity curve."""
    steps = random.randint(26, 36)
    track: list[int] = []

    for i in range(1, steps + 1):
        t = i / steps
        eased = 0.5 * (1 - math.cos(math.pi * t))
        track.append(round(distance * eased))

    deltas = [track[0]]
    for i in range(1, len(track)):
        deltas.append(track[i] - track[i - 1])

    return deltas


def solve_dingxiang_captcha(page, max_attempts: int = 7, on_status: Any = None) -> bool:
    """Automatically detect and solve Dingxiang sliding puzzle captcha."""
    bar = page.locator("[id*='dx_captcha_basic_bar']").first

    # Wait up to 5 seconds for captcha to appear
    for _ in range(25):
        if bar.is_visible():
            break
        page.wait_for_timeout(200)

    if not bar.is_visible():
        return True

    if on_status:
        on_status("Captcha detected, solving...")

    offsets = [0, 1, -1, 2, -2, 0, 1, -1]

    for attempt in range(max_attempts):
        if not bar.is_visible():
            if on_status:
                on_status("Captcha solved!")
            return True

        if on_status:
            on_status(f"Captcha attempt {attempt + 1}/{max_attempts}...")

        # Wait for canvas to draw
        page.wait_for_timeout(1500)

        canvas_data = page.evaluate("""() => {
            const canvas = document.querySelector('#dx_captcha_basic_bg_1 canvas') || document.querySelector('.dx_captcha_basic_bg canvas');
            return canvas ? canvas.toDataURL('image/png') : null;
        }""")
        margin_top_str = page.evaluate("""() => {
            const sub = document.querySelector('[id*="dx_captcha_basic_sub-slider"]');
            return sub ? sub.style.marginTop : '0px';
        }""")

        if not canvas_data or not margin_top_str:
            page.wait_for_timeout(1000)
            continue

        margin_top = float(margin_top_str.replace("px", "").strip() or 0)
        bar_box = bar.bounding_box()
        if not bar_box:
            page.wait_for_timeout(1000)
            continue

        try:
            raw_bytes = base64.b64decode(canvas_data.split(",")[1])
            hole_x_288 = find_hole_x_in_288(raw_bytes, margin_top)
            offset = offsets[attempt % len(offsets)]
            drag_distance = int(round(hole_x_288 - 21)) + offset
        except Exception:
            page.wait_for_timeout(1000)
            continue

        start_x = bar_box["x"] + 22
        start_y = bar_box["y"] + 15

        page.mouse.move(start_x, start_y)
        time.sleep(random.uniform(0.1, 0.18))
        page.mouse.down()
        time.sleep(random.uniform(0.1, 0.15))

        deltas = generate_natural_track(drag_distance)
        cur_x = start_x
        total_steps = len(deltas)
        for step_idx, dx in enumerate(deltas):
            cur_x += dx
            wobble = math.sin((step_idx / total_steps) * math.pi) * random.uniform(1.0, 2.5)
            page.mouse.move(cur_x, start_y + wobble)
            time.sleep(random.uniform(0.012, 0.024))

        time.sleep(random.uniform(0.18, 0.28))
        page.mouse.up()

        # Wait for verify response
        page.wait_for_timeout(3500)

        if not bar.is_visible():
            if on_status:
                on_status("Captcha verified!")
            return True

        # Failed on this attempt, click refresh button to get a clean puzzle
        refresh_btn = page.locator("[id*='btn-refresh']").first
        if refresh_btn.is_visible():
            try:
                refresh_btn.click()
            except Exception:
                pass
            # Wait for canvas to load the new puzzle
            for _ in range(25):
                page.wait_for_timeout(200)
                new_data = page.evaluate("""() => {
                    const canvas = document.querySelector('#dx_captcha_basic_bg_1 canvas') || document.querySelector('.dx_captcha_basic_bg canvas');
                    return canvas ? canvas.toDataURL('image/png') : null;
                }""")
                if new_data and new_data != canvas_data:
                    break
        page.wait_for_timeout(1000)

    return not bar.is_visible()


def check_oppo_serial(page, serial: str, on_status: Any = None) -> dict[str, Any]:
    """Submit a single serial number and extract device verification information."""
    captured_payload: dict[str, Any] = {}

    def on_response(response):
        if "getdeviceinfo" in response.url.lower():
            try:
                captured_payload["data"] = response.json()
            except Exception:
                pass

    page.on("response", on_response)

    try:
        if on_status:
            on_status("Loading portal...")
        page.goto("https://support.oppo.com/ua/check/", wait_until="domcontentloaded", timeout=40000)
        page.wait_for_timeout(2500)

        input_box = page.locator("input.el-input__inner").first
        input_box.fill(serial)
        page.wait_for_timeout(400)

        if on_status:
            on_status("Submitting serial...")
        search_btn = page.locator(".searchBtn").first
        search_btn.click()
        page.wait_for_timeout(2000)

        captcha_solved = solve_dingxiang_captcha(page, on_status=on_status)

        if not captcha_solved:
            return {
                "serial": serial,
                "is_active": "Unknown",
                "activation_date": "",
                "activation_comment": "",
                "model_name": "",
                "specs": "",
                "color": "",
                "certified_model": "",
                "region": "",
                "warranty_end": "",
                "status": "Captcha Failed",
                "raw_msg": "Captcha could not be resolved automatically",
            }

        # Wait for API or DOM to update
        page.wait_for_timeout(4000)

        # Extract live DOM text from the result card for activation time
        dom_val = ""
        try:
            dom_val = page.evaluate("""() => {
                const items = Array.from(document.querySelectorAll('li, div.flex, .result-item'));
                for (let item of items) {
                    const label = item.querySelector('.label');
                    const val = item.querySelector('.value');
                    if (label && val && (label.innerText || '').toLowerCase().includes('активац')) {
                        return (val.innerText || '').trim();
                    }
                }
                const all = Array.from(document.querySelectorAll('*'));
                for (let el of all) {
                    const txt = (el.innerText || '');
                    if (txt.includes('Час активації：') || txt.includes('Час активації:')) {
                        const parts = txt.split(/Час активації[：:]/);
                        if (parts.length > 1) {
                            return parts[1].split('\\n')[0].trim();
                        }
                    }
                }
                return '';
            }""") or ""
        except Exception:
            dom_val = ""

        # Check API response first
        api_data = captured_payload.get("data")
        if isinstance(api_data, dict):
            code = str(api_data.get("code", ""))
            d = api_data.get("data") or {}

            if code == "1" and d:
                reg_date = d.get("regDate")
                pmsg = str(d.get("pmsg") or "")

                dom_lower = dom_val.lower()
                pmsg_lower = pmsg.lower()

                if (
                    "передано до системи" in dom_lower
                    or "7 днів" in dom_lower
                    or "7 days" in dom_lower
                    or "transferred" in dom_lower
                    or "上报系统" in pmsg_lower
                    or "7天" in pmsg_lower
                ):
                    is_active = "Pending"
                    act_date = "Pending Sync"
                    act_comment = "The activation time has been transferred to the system. Please check again after 7 days."
                elif (
                    "не активовано" in dom_lower
                    or "not activat" in dom_lower
                    or "注册时间为空" in pmsg_lower
                    or (not reg_date and not dom_val)
                ):
                    is_active = "No"
                    act_date = "Not Activated"
                    act_comment = "Not Activated"
                elif reg_date:
                    is_active = "Yes"
                    act_date = str(reg_date)
                    act_comment = str(reg_date)
                elif dom_val:
                    is_active = "Yes"
                    act_date = dom_val
                    act_comment = dom_val
                else:
                    is_active = "No"
                    act_date = "Not Activated"
                    act_comment = pmsg or "Not Activated"

                ram = str(d.get("productRam") or "").strip()
                rom = str(d.get("productRom") or "").strip()
                specs = f"{ram}+{rom}" if ram and rom else (ram or rom)

                return {
                    "serial": serial,
                    "is_active": is_active,
                    "activation_date": act_date,
                    "activation_comment": act_comment,
                    "model_name": str(d.get("marketingModelName") or "").strip(),
                    "specs": specs,
                    "color": str(d.get("productColor") or "").strip(),
                    "certified_model": str(d.get("certifiedModel") or "").strip(),
                    "region": str(d.get("region") or "").strip(),
                    "warranty_end": str(d.get("warrantyEndTime") or "").strip(),
                    "status": "Found",
                    "raw_msg": pmsg,
                }
            elif code != "1":
                msg = str(api_data.get("msg") or "Device not found")
                return {
                    "serial": serial,
                    "is_active": "Not Found",
                    "activation_date": "",
                    "activation_comment": "",
                    "model_name": "",
                    "specs": "",
                    "color": "",
                    "certified_model": "",
                    "region": "",
                    "warranty_end": "",
                    "status": "Not Found",
                    "raw_msg": msg,
                }

        # Fallback to DOM inspection
        text = page.inner_text("body")
        if "Результат пошуку" in text:
            dom_lower = dom_val.lower()
            if "передано до системи" in dom_lower or "7 днів" in dom_lower or "7 days" in dom_lower:
                is_active = "Pending"
                act_date = "Pending Sync"
                act_comment = "The activation time has been transferred to the system. Please check again after 7 days."
            elif "не активовано" in dom_lower or "не активовано" in text.lower():
                is_active = "No"
                act_date = "Not Activated"
                act_comment = "Not Activated"
            elif dom_val:
                is_active = "Yes"
                act_date = dom_val
                act_comment = dom_val
            else:
                is_active = "Yes"
                act_date = "Activated"
                act_comment = "Activated"

            return {
                "serial": serial,
                "is_active": is_active,
                "activation_date": act_date,
                "activation_comment": act_comment,
                "model_name": "",
                "specs": "",
                "color": "",
                "certified_model": "",
                "region": "",
                "warranty_end": "",
                "status": "Found (DOM fallback)",
                "raw_msg": "Extracted from page text",
            }
        elif "не знайдено" in text.lower():
            return {
                "serial": serial,
                "is_active": "Not Found",
                "activation_date": "",
                "activation_comment": "",
                "model_name": "",
                "specs": "",
                "color": "",
                "certified_model": "",
                "region": "",
                "warranty_end": "",
                "status": "Not Found",
                "raw_msg": "Device not found in system",
            }

        return {
            "serial": serial,
            "is_active": "Unknown",
            "activation_date": "",
            "activation_comment": "",
            "model_name": "",
            "specs": "",
            "color": "",
            "certified_model": "",
            "region": "",
            "warranty_end": "",
            "status": "No Response",
            "raw_msg": "Neither API response nor result card found",
        }
    except Exception as exc:
        return {
            "serial": serial,
            "is_active": "Error",
            "activation_date": "",
            "activation_comment": "",
            "model_name": "",
            "specs": "",
            "color": "",
            "certified_model": "",
            "region": "",
            "warranty_end": "",
            "status": "Error",
            "raw_msg": str(exc),
        }
    finally:
        try:
            page.remove_listener("response", on_response)
        except Exception:
            pass


def save_to_excel(results: list[dict[str, Any]], output_path: Path, quiet: bool = False) -> None:
    """Save or update results in Excel spreadsheet with clean formatting."""
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # If file exists and openpyxl is available, read existing data to allow incremental updating
    existing_rows: dict[str, dict[str, Any]] = {}

    columns = [
        ("serial", "Serial / IMEI"),
        ("is_active", "Active?"),
        ("activation_date", "Activation Date"),
        ("activation_comment", "Activation Comment"),
        ("model_name", "Model Name"),
        ("specs", "RAM/ROM"),
        ("color", "Color"),
        ("certified_model", "Certified Model"),
        ("region", "Region"),
        ("warranty_end", "Warranty End"),
        ("status", "Check Status"),
        ("checked_at", "Last Checked"),
    ]

    if output_path.exists() and HAS_OPENPYXL:
        try:
            wb_old = openpyxl.load_workbook(output_path)
            ws_old = wb_old.active
            headers = [cell.value for cell in ws_old[1]]
            header_map = {col_label: col_key for col_key, col_label in columns}

            for row in ws_old.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                row_dict = {}
                for idx, cell_val in enumerate(row):
                    if idx < len(headers) and headers[idx] in header_map:
                        row_dict[header_map[headers[idx]]] = cell_val
                sn = str(row[0]).strip()
                if sn:
                    if not row_dict.get("activation_comment"):
                        act_d = str(row_dict.get("activation_date") or "").strip()
                        if act_d and act_d != "Not Activated":
                            row_dict["activation_comment"] = act_d
                        elif row_dict.get("is_active") == "No":
                            row_dict["activation_comment"] = "Not Activated"
                    existing_rows[sn] = row_dict
        except Exception as e:
            if not quiet:
                print(f"[Notice] Could not read existing Excel file: {e}")

    # Merge new results into existing rows without downgrading successful ones
    for r in results:
        sn = str(r["serial"]).strip()
        new_status = str(r.get("status") or "")
        old_status = str(existing_rows.get(sn, {}).get("status") or "")
        if sn in existing_rows and old_status == "Found" and new_status != "Found":
            continue
        row_entry = {**r, "checked_at": timestamp}
        existing_rows[sn] = row_entry

    final_rows = list(existing_rows.values())

    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "OPPO Status"

        # Headers
        header_fill = PatternFill(start_color="107C41", end_color="107C41", fill_type="solid")  # OPPO green
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_side = Side(border_style="thin", color="D9D9D9")
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[1].height = 28

        for col_idx, (_, col_title) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Rows
        for row_idx, data_row in enumerate(final_rows, start=2):
            ws.row_dimensions[row_idx].height = 22
            is_active_val = data_row.get("is_active", "")

            # Highlight based on active status
            status_fill = None
            if is_active_val == "Yes":
                status_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")  # Light green
            elif is_active_val == "No":
                status_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Light yellow
            elif is_active_val in ["Pending", "Syncing"]:
                status_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")  # Light blue
            elif is_active_val in ["Not Found", "Error", "Captcha Failed"]:
                status_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # Light red

            for col_idx, (col_key, _) in enumerate(columns, start=1):
                val = data_row.get(col_key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = cell_border

                if col_key in ["serial", "model_name", "color", "activation_comment"]:
                    cell.alignment = left_align
                else:
                    cell.alignment = center_align

                if col_key == "is_active" and status_fill:
                    cell.fill = status_fill
                    cell.font = Font(name="Calibri", size=10, bold=True)

        # Auto-adjust column widths
        for col_idx in range(1, len(columns) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, len(final_rows) + 2))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        wb.save(output_path)
        if not quiet:
            print(f"Successfully saved {len(final_rows)} records to Excel: {output_path}")
    elif HAS_PANDAS:
        df = pd.DataFrame(final_rows)
        df.to_excel(output_path, index=False)
        if not quiet:
            print(f"Saved {len(final_rows)} records to Excel: {output_path}")
    else:
        # Fallback to CSV
        csv_path = output_path.with_suffix(".csv")
        import csv
        with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=[k for k, _ in columns], extrasaction="ignore")
            writer.writeheader()
            writer.writerows(final_rows)
        if not quiet:
            print(f"Saved to CSV (openpyxl not available): {csv_path}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Check OPPO serial/IMEI activation status and update Excel.")
    parser.add_argument("--input", type=Path, default=Path("serials.txt"), help="Path to input serials file (default: serials.txt)")
    parser.add_argument("--output", type=Path, default=Path("oppo_results.xlsx"), help="Path to output Excel file (default: oppo_results.xlsx)")
    parser.add_argument("--headful", action="store_true", help="Launch visible browser window instead of headless mode")
    parser.add_argument("--delay", type=float, default=2.0, help="Delay between checks in seconds (default: 2.0)")
    parser.add_argument("--skip-found", action="store_true", help="Skip serials that are already successfully verified in output Excel")
    parser.add_argument("--limit", type=int, default=None, help="Limit number of serials to process (optional)")
    args = parser.parse_args()

    identifiers = read_identifiers(args.input)
    if args.limit:
        identifiers = identifiers[: args.limit]

    if args.skip_found and args.output.exists() and HAS_OPENPYXL:
        try:
            wb_check = openpyxl.load_workbook(args.output)
            ws_check = wb_check.active
            headers = [cell.value for cell in ws_check[1]]
            status_idx = headers.index("Check Status") if "Check Status" in headers else -1
            sn_idx = headers.index("Serial / IMEI") if "Serial / IMEI" in headers else 0
            found_sns = set()
            for row in ws_check.iter_rows(min_row=2, values_only=True):
                if row and status_idx != -1 and len(row) > status_idx and row[status_idx] == "Found":
                    found_sns.add(str(row[sn_idx]).strip())
            initial_count = len(identifiers)
            identifiers = [s for s in identifiers if s not in found_sns]
            if initial_count != len(identifiers):
                print(f"Skipping {initial_count - len(identifiers)} already verified serial(s). {len(identifiers)} remaining to check.")
        except Exception:
            pass

    if not identifiers:
        print(f"All serials in {args.input} are already verified in {args.output}!")
        return 0

    print(f"Loaded {len(identifiers)} serial(s) to check.")
    print(f"Target Excel output: {args.output}\n")

    results: list[dict[str, Any]] = []

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not args.headful)
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            viewport={"width": 1280, "height": 800},
        )
        page = context.new_page()

        if HAS_TQDM:
            pbar = tqdm(
                identifiers,
                desc="Checking OPPO Devices",
                unit="device",
                ncols=95,
                bar_format="{l_bar}{bar}| {n_fmt}/{total_fmt} [{elapsed}<{remaining}] {postfix}",
            )
            for idx, sn in enumerate(pbar, start=1):
                pbar.set_postfix_str(f"{sn}: Connecting...")
                res = check_oppo_serial(page, sn, on_status=lambda msg: pbar.set_postfix_str(f"{sn}: {msg}"))
                results.append(res)

                # Save incrementally so no data is lost
                save_to_excel(results, args.output, quiet=True)

                status_line = f"[{idx}/{len(identifiers)}] {sn} -> Active: {res['is_active']} | Model: {res['model_name']} | Color: {res['color']}"
                tqdm.write(status_line)
                pbar.set_postfix_str(f"Last: {res['is_active']}")

                if idx < len(identifiers):
                    time.sleep(args.delay)
            pbar.close()
        else:
            for idx, sn in enumerate(identifiers, start=1):
                pct = int((idx / len(identifiers)) * 100)
                bar_str = "[" + "=" * (pct // 5) + " " * (20 - pct // 5) + "]"
                print(f"{bar_str} {pct}% ({idx}/{len(identifiers)}) Checking: {sn} ...", flush=True)

                res = check_oppo_serial(page, sn, on_status=lambda msg: print(f"   [{sn}] {msg}", flush=True))
                results.append(res)
                save_to_excel(results, args.output, quiet=True)

                print(f"   Result: Active={res['is_active']} | Model={res['model_name']} | Color={res['color']}", flush=True)
                if idx < len(identifiers):
                    time.sleep(args.delay)

        browser.close()

    # Final save and summary
    save_to_excel(results, args.output, quiet=False)
    print("\n" + "=" * 60)
    print(f"Done! Checked {len(results)} device(s). Results written to {args.output}")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""SQLite Database Manager for OPPO Device Verification & Sales Activation Monitor.

Thread-safe SQLite connection manager with WAL mode, device tracking,
group segmentation, run history, and dynamic user settings.
"""

from __future__ import annotations

import csv
import datetime
import hashlib
import os
import secrets
import sqlite3
import threading
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Generator

import openpyxl

DB_DIR = Path(__file__).resolve().parent / "data"
DB_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DB_DIR / "oppo_monitor.db"

_db_lock = threading.Lock()


def get_db_path() -> Path:
    return DB_PATH


@contextmanager
def get_db_cursor() -> Generator[sqlite3.Cursor, None, None]:
    """Provide a thread-safe SQLite connection cursor with Row factory and WAL mode."""
    with _db_lock:
        conn = sqlite3.connect(
            str(DB_PATH),
            timeout=30.0,
            check_same_thread=False,
        )
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL;")
        conn.execute("PRAGMA synchronous=NORMAL;")
        conn.execute("PRAGMA foreign_keys=ON;")
        cursor = conn.cursor()
        try:
            yield cursor
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()


def init_db() -> None:
    """Initialize database tables, indexes, and default settings."""
    with get_db_cursor() as cur:
        # Devices table
        cur.execute("""
            CREATE TABLE IF NOT EXISTS devices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                serial TEXT UNIQUE NOT NULL,
                group_name TEXT NOT NULL DEFAULT 'Default',
                is_active TEXT NOT NULL DEFAULT 'Unknown',
                activation_date TEXT DEFAULT '',
                activation_comment TEXT DEFAULT '',
                model_name TEXT DEFAULT '',
                specs TEXT DEFAULT '',
                color TEXT DEFAULT '',
                certified_model TEXT DEFAULT '',
                region TEXT DEFAULT '',
                warranty_end TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT 'Pending',
                raw_msg TEXT DEFAULT '',
                first_added_at TEXT NOT NULL,
                last_checked_at TEXT DEFAULT '',
                activated_at TEXT DEFAULT '',
                notified_activation INTEGER NOT NULL DEFAULT 0,
                notes TEXT DEFAULT ''
            );
        """)

        # Indexes for fast querying & filtering
        cur.execute("CREATE INDEX IF NOT EXISTS idx_devices_serial ON devices(serial);")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_devices_group ON devices(group_name);")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_devices_is_active ON devices(is_active);")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_devices_status ON devices(status);")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_devices_activated_at ON devices(activated_at);")

        # Groups table
        cur.execute("""
            CREATE TABLE IF NOT EXISTS groups (
                name TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                description TEXT DEFAULT ''
            );
        """)
        # Insert Default and Not Run groups
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cur.execute("INSERT OR IGNORE INTO groups (name, created_at, description) VALUES (?, ?, ?);", ("Default", now, "Standard default group"))
        cur.execute("INSERT OR IGNORE INTO groups (name, created_at, description) VALUES (?, ?, ?);", ("Not Run", now, "New serials awaiting initial verification check"))

        # Run history table
        cur.execute("""
            CREATE TABLE IF NOT EXISTS run_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_type TEXT NOT NULL,
                target_group TEXT DEFAULT '',
                started_at TEXT NOT NULL,
                finished_at TEXT DEFAULT '',
                total_checked INTEGER NOT NULL DEFAULT 0,
                newly_activated_count INTEGER NOT NULL DEFAULT 0,
                full_sheet_path TEXT DEFAULT '',
                newly_activated_sheet_path TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT 'running',
                summary_json TEXT DEFAULT '{}'
            );
        """)

        # Settings table
        cur.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL,
                description TEXT DEFAULT ''
            );
        """)

        # Default configuration settings
        default_settings = [
            ("schedule_interval_minutes", "60", "How often to check unactivated serials automatically (in minutes, 0=disabled)"),
            ("schedule_enabled", "1", "1 if auto-scheduler is enabled, 0 if paused"),
            ("threads_count", "2", "Number of parallel browser worker threads (1-4)"),
            ("delay_seconds", "2.0", "Delay between serial checks in seconds"),
            ("fine_tune_offset", "0.0", "Slider offset adjust in pixels"),
            ("headful_mode", "0", "Launch visible browser windows on server (0=headless, 1=visible)"),
            ("auto_recover_failures", "1", "Automatically run recovery pass for errors at end of batch"),
            ("telegram_bot_token", "", "Telegram Bot Token from @BotFather"),
            ("telegram_chat_id", "", "Telegram Chat ID or Group ID to send alerts to"),
            ("telegram_notifications_enabled", "0", "1 to enable instant Telegram alerts, 0 to disable"),
            ("telegram_send_sheets", "1", "1 to automatically attach both Excel sheets to Telegram upon run completion, 0 to disable"),
            ("telegram_notify_runs", "1", "1 to send Telegram alerts when verification runs start & finish, 0 to disable"),
        ]

        for k, v, d in default_settings:
            cur.execute("""
                INSERT OR IGNORE INTO settings (key, value, description)
                VALUES (?, ?, ?);
            """, (k, v, d))

        # Users table
        cur.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                salt TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'customer',
                created_at TEXT NOT NULL
            );
        """)

        # Sessions table
        cur.execute("""
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                username TEXT NOT NULL,
                role TEXT NOT NULL,
                created_at TEXT NOT NULL,
                expires_at TEXT NOT NULL
            );
        """)

    # Seed default user accounts
    ensure_default_users()


# --- Authentication & User Management ---

def hash_password(password: str, salt: str | None = None) -> tuple[str, str]:
    """Generate PBKDF2-SHA256 password hash and salt."""
    if not salt:
        salt = secrets.token_hex(16)
    hashed = hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), salt.encode("utf-8"), 100000
    ).hex()
    return hashed, salt


def verify_password(password: str, stored_hash: str, salt: str) -> bool:
    """Verify password against stored PBKDF2-SHA256 hash."""
    hashed, _ = hash_password(password, salt)
    return secrets.compare_digest(hashed, stored_hash)


def ensure_default_users() -> None:
    """Ensure the required users ('eslam' admin and 'customer') exist with correct credentials."""
    users_to_seed = [
        ("eslam", "Solom_2040", "admin"),
        ("customer", "1234", "customer"),
    ]
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    with get_db_cursor() as cur:
        for username, password, role in users_to_seed:
            pwd_hash, salt = hash_password(password)
            cur.execute("SELECT id FROM users WHERE username = ?;", (username,))
            row = cur.fetchone()
            if row:
                cur.execute(
                    """
                    UPDATE users 
                    SET password_hash = ?, salt = ?, role = ? 
                    WHERE username = ?;
                    """,
                    (pwd_hash, salt, role, username),
                )
            else:
                cur.execute(
                    """
                    INSERT INTO users (username, password_hash, salt, role, created_at)
                    VALUES (?, ?, ?, ?, ?);
                    """,
                    (username, pwd_hash, salt, role, now),
                )


def authenticate_user(username: str, password: str) -> dict[str, Any] | None:
    """Authenticate username and password, returning user dict if successful."""
    username = username.strip().lower()
    with get_db_cursor() as cur:
        cur.execute(
            "SELECT id, username, password_hash, salt, role FROM users WHERE LOWER(username) = ?;",
            (username,),
        )
        row = cur.fetchone()
        if not row:
            return None
        if verify_password(password, row["password_hash"], row["salt"]):
            return {
                "id": row["id"],
                "username": row["username"],
                "role": row["role"],
            }
        return None


def create_session(username: str, role: str, duration_days: int = 30) -> str:
    """Create a new authenticated session token and record in database."""
    token = secrets.token_urlsafe(32)
    now_dt = datetime.datetime.now()
    created_at = now_dt.strftime("%Y-%m-%d %H:%M:%S")
    expires_at = (now_dt + datetime.timedelta(days=duration_days)).strftime("%Y-%m-%d %H:%M:%S")

    with get_db_cursor() as cur:
        cur.execute(
            """
            INSERT INTO sessions (token, username, role, created_at, expires_at)
            VALUES (?, ?, ?, ?, ?);
            """,
            (token, username, role, created_at, expires_at),
        )
    return token


def get_session(token: str) -> dict[str, Any] | None:
    """Retrieve active session by token, validating expiration."""
    if not token:
        return None
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_db_cursor() as cur:
        cur.execute(
            """
            SELECT username, role, expires_at 
            FROM sessions 
            WHERE token = ? AND expires_at > ?;
            """,
            (token, now_str),
        )
        row = cur.fetchone()
        if not row:
            return None
        return {
            "username": row["username"],
            "role": row["role"],
        }


def delete_session(token: str) -> None:
    """Invalidate and remove session token."""
    if not token:
        return
    with get_db_cursor() as cur:
        cur.execute("DELETE FROM sessions WHERE token = ?;", (token,))


# --- Settings Helpers ---

def get_setting(key: str, default: str = "") -> str:
    """Retrieve setting value by key."""
    with get_db_cursor() as cur:
        cur.execute("SELECT value FROM settings WHERE key = ?;", (key,))
        row = cur.fetchone()
        return str(row["value"]) if row else default


def get_all_settings() -> dict[str, str]:
    """Retrieve all system settings as a dictionary."""
    with get_db_cursor() as cur:
        cur.execute("SELECT key, value FROM settings;")
        return {str(r["key"]): str(r["value"]) for r in cur.fetchall()}


def set_setting(key: str, value: str) -> None:
    """Update or insert a system setting."""
    with get_db_cursor() as cur:
        cur.execute("""
            INSERT INTO settings (key, value)
            VALUES (?, ?)
            ON CONFLICT(key) DO UPDATE SET value = excluded.value;
        """, (key, str(value)))


def set_multiple_settings(settings_dict: dict[str, Any]) -> None:
    """Bulk update multiple settings."""
    with get_db_cursor() as cur:
        for k, v in settings_dict.items():
            cur.execute("""
                INSERT INTO settings (key, value)
                VALUES (?, ?)
                ON CONFLICT(key) DO UPDATE SET value = excluded.value;
            """, (k, str(v)))


# --- Group Helpers ---

def get_groups() -> list[dict[str, Any]]:
    """Get all groups along with counts of total and active devices."""
    with get_db_cursor() as cur:
        cur.execute("""
            SELECT 
                g.name,
                g.created_at,
                g.description,
                COUNT(d.id) AS total_devices,
                SUM(CASE WHEN d.is_active = 'Yes' THEN 1 ELSE 0 END) AS active_devices,
                SUM(CASE WHEN d.is_active != 'Yes' OR d.is_active IS NULL THEN 1 ELSE 0 END) AS pending_devices
            FROM groups g
            LEFT JOIN devices d ON g.name = d.group_name
            GROUP BY g.name
            ORDER BY g.name ASC;
        """)
        return [dict(r) for r in cur.fetchall()]


def create_group(name: str, description: str = "") -> bool:
    """Create a new device group."""
    name = name.strip()
    if not name:
        return False
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_db_cursor() as cur:
        cur.execute("""
            INSERT OR IGNORE INTO groups (name, created_at, description)
            VALUES (?, ?, ?);
        """, (name, now, description.strip()))
        return cur.rowcount > 0


def delete_group(name: str) -> bool:
    """Delete a group and reassign its devices to 'Default'."""
    name = name.strip()
    if name.lower() == "default":
        return False
    with get_db_cursor() as cur:
        cur.execute("UPDATE devices SET group_name = 'Default' WHERE group_name = ?;", (name,))
        cur.execute("DELETE FROM groups WHERE name = ?;", (name,))
        return cur.rowcount > 0


# --- Device Helpers ---

def upsert_device(serial: str, group_name: str = "Not Run", notes: str = "") -> dict[str, Any]:
    """Add a new serial or update its group."""
    serial = str(serial).strip().upper()
    group_name = str(group_name).strip() or "Not Run"
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    with get_db_cursor() as cur:
        # Ensure group exists
        cur.execute("INSERT OR IGNORE INTO groups (name, created_at) VALUES (?, ?);", (group_name, now))

        cur.execute("""
            INSERT INTO devices (serial, group_name, first_added_at, is_active, status, notes)
            VALUES (?, ?, ?, 'Not Run', 'Not Run', ?)
            ON CONFLICT(serial) DO UPDATE SET
                group_name = CASE WHEN excluded.group_name != 'Not Run' THEN excluded.group_name ELSE devices.group_name END,
                notes = CASE WHEN excluded.notes != '' THEN excluded.notes ELSE devices.notes END;
        """, (serial, group_name, now, notes))

        cur.execute("SELECT * FROM devices WHERE serial = ?;", (serial,))
        row = cur.fetchone()
        return dict(row) if row else {}


def bulk_upsert_devices(serials: list[str], group_name: str = "Not Run") -> tuple[int, int]:
    """Bulk insert serials into a group. Returns (added_count, total_provided)."""
    group_name = str(group_name).strip() or "Not Run"
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    added = 0
    clean_serials = [s.strip().upper() for s in serials if s and s.strip()]

    with get_db_cursor() as cur:
        cur.execute("INSERT OR IGNORE INTO groups (name, created_at) VALUES (?, ?);", (group_name, now))
        for sn in clean_serials:
            cur.execute("""
                INSERT INTO devices (serial, group_name, first_added_at, is_active, status)
                VALUES (?, ?, ?, 'Not Run', 'Not Run')
                ON CONFLICT(serial) DO NOTHING;
            """, (sn, group_name, now))
            if cur.rowcount > 0:
                added += 1

    return added, len(clean_serials)


def update_device_result(serial: str, res: dict[str, Any]) -> tuple[dict[str, Any], bool]:
    """
    Update device verification result in database.
    Detects if the device transitioned from not active/unknown -> Active ('Yes').
    Returns (updated_device_dict, is_newly_activated).
    """
    serial = str(serial).strip().upper()
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    with get_db_cursor() as cur:
        cur.execute("SELECT * FROM devices WHERE serial = ?;", (serial,))
        existing = cur.fetchone()

        prev_is_active = existing["is_active"] if existing else "Unknown"
        new_is_active = str(res.get("is_active") or "Unknown").strip()

        # Activation transition detection:
        # If it was NOT active before (or Unknown/Pending) and now is "Yes"
        is_newly_activated = (prev_is_active != "Yes" and new_is_active == "Yes")

        activated_at = existing["activated_at"] if existing and existing["activated_at"] else ""
        if is_newly_activated:
            activated_at = now

        cur.execute("""
            UPDATE devices SET
                is_active = ?,
                activation_date = ?,
                activation_comment = ?,
                model_name = ?,
                specs = ?,
                color = ?,
                certified_model = ?,
                region = ?,
                warranty_end = ?,
                status = ?,
                raw_msg = ?,
                last_checked_at = ?,
                activated_at = ?
            WHERE serial = ?;
        """, (
            new_is_active,
            str(res.get("activation_date") or ""),
            str(res.get("activation_comment") or ""),
            str(res.get("model_name") or ""),
            str(res.get("specs") or ""),
            str(res.get("color") or ""),
            str(res.get("certified_model") or ""),
            str(res.get("region") or ""),
            str(res.get("warranty_end") or ""),
            str(res.get("status") or ""),
            str(res.get("raw_msg") or ""),
            now,
            activated_at,
            serial,
        ))

        cur.execute("SELECT * FROM devices WHERE serial = ?;", (serial,))
        updated_row = cur.fetchone()
        return dict(updated_row) if updated_row else {}, is_newly_activated


def mark_device_notified(serial: str) -> None:
    """Mark device as having had its activation notification sent."""
    with get_db_cursor() as cur:
        cur.execute("UPDATE devices SET notified_activation = 1 WHERE serial = ?;", (serial.strip().upper(),))


def get_device_filter_options() -> dict[str, Any]:
    """Retrieve distinct models, colors, specs, and groups with counts for filtering UI."""
    with get_db_cursor() as cur:
        cur.execute("""
            SELECT model_name, COUNT(*) as count 
            FROM devices 
            WHERE model_name IS NOT NULL AND model_name != '' 
            GROUP BY model_name 
            ORDER BY count DESC, model_name ASC;
        """)
        models = [{"name": r["model_name"], "count": r["count"]} for r in cur.fetchall()]

        cur.execute("""
            SELECT color, COUNT(*) as count 
            FROM devices 
            WHERE color IS NOT NULL AND color != '' 
            GROUP BY color 
            ORDER BY count DESC, color ASC;
        """)
        colors = [{"name": r["color"], "count": r["count"]} for r in cur.fetchall()]

        cur.execute("""
            SELECT specs, COUNT(*) as count 
            FROM devices 
            WHERE specs IS NOT NULL AND specs != '' 
            GROUP BY specs 
            ORDER BY count DESC, specs ASC;
        """)
        specs = [{"name": r["specs"], "count": r["count"]} for r in cur.fetchall()]

        cur.execute("""
            SELECT g.name, 
                   COUNT(d.id) as total_devices,
                   SUM(CASE WHEN d.is_active = 'Yes' THEN 1 ELSE 0 END) as active_devices
            FROM groups g
            LEFT JOIN devices d ON g.name = d.group_name
            GROUP BY g.name
            ORDER BY g.name ASC;
        """)
        groups = [{"name": r["name"], "total": r["total_devices"], "active": r["active_devices"]} for r in cur.fetchall()]

        cur.execute("""
            SELECT activation_date, COUNT(*) as count 
            FROM devices 
            WHERE is_active = 'Yes' 
              AND activation_date IS NOT NULL 
              AND activation_date != '' 
              AND activation_date NOT IN ('Not Activated', 'Pending Sync', 'Unknown')
            GROUP BY activation_date 
            ORDER BY count DESC;
        """)
        activation_dates = [{"name": r["activation_date"], "count": r["count"]} for r in cur.fetchall()]

        return {
            "models": models,
            "colors": colors,
            "specs": specs,
            "groups": groups,
            "activation_dates": activation_dates,
        }


def get_devices(
    search: str = "",
    group_name: str = "",
    status_filter: str = "",
    model: str = "",
    color: str = "",
    specs: str = "",
    activation_date: str = "",
    newly_activated_only: bool = False,
    limit: int = 100,
    offset: int = 0,
) -> tuple[list[dict[str, Any]], int]:
    """Retrieve devices with filtering, pagination, and total count."""
    query = "SELECT * FROM devices WHERE 1=1"
    count_query = "SELECT COUNT(*) AS cnt FROM devices WHERE 1=1"
    params: list[Any] = []

    if search.strip():
        # Split search into terms so e.g. "A6x purple" matches across model and color
        terms = [t.strip() for t in search.strip().split() if t.strip()]
        for term in terms:
            term_wildcard = f"%{term}%"
            term_clause = """ AND (
                serial LIKE ? OR 
                model_name LIKE ? OR 
                color LIKE ? OR 
                specs LIKE ? OR 
                certified_model LIKE ? OR 
                region LIKE ? OR 
                notes LIKE ? OR
                activation_date LIKE ?
            )"""
            query += term_clause
            count_query += term_clause
            params.extend([term_wildcard] * 8)

    if group_name.strip() and group_name.strip().lower() != "all":
        query += " AND group_name = ?"
        count_query += " AND group_name = ?"
        params.append(group_name.strip())

    if model.strip() and model.strip().lower() != "all":
        query += " AND model_name = ?"
        count_query += " AND model_name = ?"
        params.append(model.strip())

    if color.strip() and color.strip().lower() != "all":
        query += " AND color = ?"
        count_query += " AND color = ?"
        params.append(color.strip())

    if specs.strip() and specs.strip().lower() != "all":
        query += " AND specs = ?"
        count_query += " AND specs = ?"
        params.append(specs.strip())

    if activation_date.strip() and activation_date.strip().lower() != "all":
        ad = activation_date.strip()
        parts = ad.split("-")
        if len(parts) == 3 and len(parts[0]) == 4:
            # Convert YYYY-MM-DD (from datepicker) to DD-MM-YYYY
            alt_ad = f"{parts[2]}-{parts[1]}-{parts[0]}"
            query += " AND (activation_date LIKE ? OR activation_date LIKE ?)"
            count_query += " AND (activation_date LIKE ? OR activation_date LIKE ?)"
            params.extend([f"%{ad}%", f"%{alt_ad}%"])
        else:
            query += " AND activation_date LIKE ?"
            count_query += " AND activation_date LIKE ?"
            params.append(f"%{ad}%")

    if status_filter.strip() and status_filter.strip().lower() != "all":
        stat = status_filter.strip()
        if stat == "Active":
            query += " AND is_active = 'Yes'"
            count_query += " AND is_active = 'Yes'"
        elif stat == "Inactive":
            query += " AND is_active = 'No'"
            count_query += " AND is_active = 'No'"
        elif stat == "Pending":
            query += " AND (is_active = 'Pending' OR status = 'Pending') AND is_active != 'Not Run'"
            count_query += " AND (is_active = 'Pending' OR status = 'Pending') AND is_active != 'Not Run'"
        elif stat == "Error":
            query += " AND (is_active IN ('Error', 'Unknown') OR status IN ('Error', 'Captcha Failed', 'No Response')) AND is_active != 'Not Run'"
            count_query += " AND (is_active IN ('Error', 'Unknown') OR status IN ('Error', 'Captcha Failed', 'No Response')) AND is_active != 'Not Run'"
        elif stat == "Not Run":
            query += " AND (is_active = 'Not Run' OR status = 'Not Run' OR last_checked_at = '')"
            count_query += " AND (is_active = 'Not Run' OR status = 'Not Run' OR last_checked_at = '')"

    if newly_activated_only:
        # Detected active today
        today_start = datetime.datetime.now().strftime("%Y-%m-%d 00:00:00")
        today_end = datetime.datetime.now().strftime("%Y-%m-%d 23:59:59")
        query += " AND is_active = 'Yes' AND activated_at >= ? AND activated_at <= ?"
        count_query += " AND is_active = 'Yes' AND activated_at >= ? AND activated_at <= ?"
        params.extend([today_start, today_end])

    query += " ORDER BY id DESC LIMIT ? OFFSET ?;"

    with get_db_cursor() as cur:
        # Get total matching count
        cur.execute(count_query, params)
        total_cnt = cur.fetchone()["cnt"]

        # Get rows
        cur.execute(query, params + [limit, offset])
        rows = [dict(r) for r in cur.fetchall()]

    return rows, total_cnt


def get_all_devices_for_export() -> list[dict[str, Any]]:
    """Retrieve all devices for Excel/CSV export."""
    with get_db_cursor() as cur:
        cur.execute("SELECT * FROM devices ORDER BY id ASC;")
        return [dict(r) for r in cur.fetchall()]


def get_serials_to_check(
    mode: str = "all",
    group_name: str = "",
    selected_serials: list[str] | None = None,
    pending_only: bool = False,
    errors_only: bool = False,
) -> list[str]:
    """Retrieve list of serials matching specific run criteria."""
    query = "SELECT serial FROM devices WHERE 1=1"
    params: list[Any] = []

    if mode == "errors" or errors_only:
        query += " AND (is_active IN ('Error', 'Unknown') OR status IN ('Error', 'Captcha Failed', 'No Response')) AND is_active != 'Not Run'"
    elif mode == "not_run":
        query += " AND (is_active = 'Not Run' OR status = 'Not Run' OR last_checked_at = '')"
    elif mode in ["pending_sync", "pending"]:
        query += " AND is_active = 'Pending'"
    elif mode in ["inactive_only", "inactive"]:
        query += " AND is_active = 'No'"
    elif mode in ["unactivated", "pending_only"] or pending_only:
        # Exclude already verified active devices
        query += " AND is_active != 'Yes'"
    elif mode == "selected" and selected_serials:
        clean = [s.strip().upper() for s in selected_serials if s.strip()]
        placeholders = ",".join("?" for _ in clean)
        query += f" AND serial IN ({placeholders})"
        params.extend(clean)
    elif mode == "group" and group_name.strip():
        query += " AND group_name = ?"
        params.append(group_name.strip())

    query += " ORDER BY id ASC;"

    with get_db_cursor() as cur:
        cur.execute(query, params)
        return [str(r["serial"]) for r in cur.fetchall()]


def delete_device(serial: str) -> bool:
    """Delete a single device by serial."""
    with get_db_cursor() as cur:
        cur.execute("DELETE FROM devices WHERE serial = ?;", (serial.strip().upper(),))
        return cur.rowcount > 0


def delete_multiple_devices(serials: list[str]) -> int:
    """Delete multiple devices by serials."""
    clean = [s.strip().upper() for s in serials if s.strip()]
    if not clean:
        return 0
    placeholders = ",".join("?" for _ in clean)
    with get_db_cursor() as cur:
        cur.execute(f"DELETE FROM devices WHERE serial IN ({placeholders});", clean)
        return cur.rowcount


def get_dashboard_stats() -> dict[str, Any]:
    """Compute high-level metrics for dashboard cards."""
    now = datetime.datetime.now()
    today_start = now.strftime("%Y-%m-%d 00:00:00")
    today_end = now.strftime("%Y-%m-%d 23:59:59")

    with get_db_cursor() as cur:
        cur.execute("""
            SELECT
                COUNT(*) AS total,
                SUM(CASE WHEN is_active = 'Yes' THEN 1 ELSE 0 END) AS active_count,
                SUM(CASE WHEN is_active = 'No' THEN 1 ELSE 0 END) AS inactive_count,
                SUM(CASE WHEN (is_active = 'Pending' OR status = 'Pending') AND is_active != 'Not Run' AND status != 'Not Run' THEN 1 ELSE 0 END) AS pending_count,
                SUM(CASE WHEN is_active = 'Not Run' OR status = 'Not Run' OR (last_checked_at = '' AND is_active NOT IN ('Yes', 'No', 'Pending')) THEN 1 ELSE 0 END) AS not_run_count,
                SUM(CASE WHEN (is_active IN ('Error', 'Unknown') OR status IN ('Error', 'Captcha Failed', 'No Response')) AND is_active != 'Not Run' AND status != 'Not Run' THEN 1 ELSE 0 END) AS error_count,
                SUM(CASE WHEN is_active = 'Yes' AND activated_at >= ? AND activated_at <= ? THEN 1 ELSE 0 END) AS activated_today_count
            FROM devices;
        """, (today_start, today_end))
        row = cur.fetchone()

        cur.execute("SELECT COUNT(*) AS cnt FROM groups;")
        groups_cnt = cur.fetchone()["cnt"]

        return {
            "total": row["total"] or 0,
            "active": row["active_count"] or 0,
            "inactive": row["inactive_count"] or 0,
            "pending": row["pending_count"] or 0,
            "not_run": row["not_run_count"] or 0,
            "errors": row["error_count"] or 0,
            "activated_today": row["activated_today_count"] or 0,
            "groups_count": groups_cnt or 0,
        }


# --- Run History Helpers ---

def create_run_history_record(run_type: str, target_group: str = "", total: int = 0) -> int:
    """Create a new run history row when a verification job starts."""
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_db_cursor() as cur:
        cur.execute("""
            INSERT INTO run_history (run_type, target_group, started_at, total_checked, status)
            VALUES (?, ?, ?, ?, 'running');
        """, (run_type, target_group, now, total))
        return cur.lastrowid or 0


def update_run_history_record(
    run_id: int,
    status: str,
    total_checked: int,
    newly_activated: int,
    full_sheet: str = "",
    newly_activated_sheet: str = "",
    summary_json: str = "{}",
) -> None:
    """Complete a run history row."""
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_db_cursor() as cur:
        cur.execute("""
            UPDATE run_history SET
                status = ?,
                finished_at = ?,
                total_checked = ?,
                newly_activated_count = ?,
                full_sheet_path = ?,
                newly_activated_sheet_path = ?,
                summary_json = ?
            WHERE id = ?;
        """, (status, now, total_checked, newly_activated, full_sheet, newly_activated_sheet, summary_json, run_id))


def get_recent_runs(limit: int = 15) -> list[dict[str, Any]]:
    """Get recent verification runs."""
    with get_db_cursor() as cur:
        cur.execute("SELECT * FROM run_history ORDER BY id DESC LIMIT ?;", (limit,))
        return [dict(r) for r in cur.fetchall()]


def import_sheet_to_database(
    file_path: Path | str = Path("oppo_results_20261017_1718.csv"),
    group_name: str = "Default",
    clear_test: bool = True,
) -> int:
    """Import devices from an existing Excel or CSV results sheet into the database."""
    p = Path(file_path).resolve()
    if not p.exists():
        # Fallback to current folder
        p = Path(__file__).resolve().parent / Path(file_path).name
        if not p.exists():
            return 0

    rows: list[dict[str, Any]] = []
    if p.suffix.lower() == ".csv":
        with open(p, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
    elif p.suffix.lower() in [".xlsx", ".xls"]:
        wb = openpyxl.load_workbook(p, data_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip() for c in ws[1]] if ws.max_row >= 1 else []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and any(r):
                row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                rows.append(row_dict)

    if not rows:
        return 0

    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    imported_count = 0

    with get_db_cursor() as cur:
        if clear_test:
            # Delete any test serials from prior tests
            cur.execute("DELETE FROM devices WHERE serial LIKE 'TEST_OPPO_%' OR serial LIKE 'API_TEST_%';")

        # Ensure group exists
        cur.execute("INSERT OR IGNORE INTO groups (name, created_at, description) VALUES (?, ?, ?);",
                    (group_name, now, "Default imported sheet group"))

        for r in rows:
            # Flexible header mapping
            sn = str(r.get("Serial / IMEI") or r.get("serial") or r.get("IMEI") or r.get("Serial") or "").strip().upper()
            if not sn or sn.lower() in ["serial", "imei", "serial / imei"]:
                continue

            is_active = str(r.get("Active?") or r.get("is_active") or "Unknown").strip()
            act_date = str(r.get("Activation Date") or r.get("activation_date") or "").strip()
            act_comment = str(r.get("Activation Comment") or r.get("activation_comment") or "").strip()
            model_name = str(r.get("Model Name") or r.get("model_name") or "").strip()
            specs = str(r.get("RAM/ROM") or r.get("specs") or r.get("RAM / Storage") or "").strip()
            color = str(r.get("Color") or r.get("color") or "").strip()
            certified = str(r.get("Certified Model") or r.get("certified_model") or "").strip()
            region = str(r.get("Region") or r.get("region") or "").strip()
            warranty = str(r.get("Warranty End") or r.get("warranty_end") or "").strip()
            status = str(r.get("Check Status") or r.get("status") or "Found").strip()
            last_checked = str(r.get("Last Checked") or r.get("last_checked_at") or now).strip()

            # Baseline historical import: leave activated_at empty.
            # Only live detection runs will stamp activated_at with the detection timestamp!
            activated_at = ""

            cur.execute("""
                INSERT INTO devices (
                    serial, group_name, is_active, activation_date, activation_comment,
                    model_name, specs, color, certified_model, region, warranty_end,
                    status, first_added_at, last_checked_at, activated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(serial) DO UPDATE SET
                    group_name = excluded.group_name,
                    is_active = excluded.is_active,
                    activation_date = excluded.activation_date,
                    activation_comment = excluded.activation_comment,
                    model_name = excluded.model_name,
                    specs = excluded.specs,
                    color = excluded.color,
                    certified_model = excluded.certified_model,
                    region = excluded.region,
                    warranty_end = excluded.warranty_end,
                    status = excluded.status,
                    last_checked_at = excluded.last_checked_at,
                    activated_at = excluded.activated_at;
            """, (
                sn, group_name, is_active, act_date, act_comment,
                model_name, specs, color, certified, region, warranty,
                status, now, last_checked, activated_at,
            ))
            imported_count += 1

    return imported_count


def get_database_info() -> dict[str, Any]:
    """Retrieve detailed database engine, storage statistics, and device metrics."""
    stat_size = DB_PATH.stat().st_size if DB_PATH.exists() else 0
    if stat_size >= 1024 * 1024:
        size_fmt = f"{stat_size / (1024 * 1024):.2f} MB"
    elif stat_size >= 1024:
        size_fmt = f"{stat_size / 1024:.1f} KB"
    else:
        size_fmt = f"{stat_size} Bytes"

    mtime_str = ""
    if DB_PATH.exists():
        mtime = datetime.datetime.fromtimestamp(DB_PATH.stat().st_mtime)
        mtime_str = mtime.strftime("%Y-%m-%d %H:%M:%S")

    stats = get_dashboard_stats()

    with get_db_cursor() as cur:
        cur.execute("SELECT COUNT(*) AS cnt FROM run_history;")
        runs_cnt = cur.fetchone()["cnt"]

    return {
        "engine": "SQLite 3 (WAL Mode)",
        "file_name": DB_PATH.name,
        "file_path": str(DB_PATH),
        "file_size": size_fmt,
        "file_size_bytes": stat_size,
        "last_modified": mtime_str,
        "total_devices": stats["total"],
        "active_devices": stats["active"],
        "inactive_devices": stats["inactive"],
        "pending_devices": stats["pending"],
        "error_devices": stats["errors"],
        "activated_today": stats["activated_today"],
        "groups_count": stats["groups_count"],
        "runs_count": runs_cnt,
    }


# Initialize on import
init_db()

"""Excel & CSV Report Generator for OPPO Device Verification & Sales Monitoring.

Generates the TWO distinct report sheets required after every run:
1. Full Run Sheet (All checked serials with full device specifications)
2. Newly Activated Sheet (Only devices detected as active in this batch)
"""

from __future__ import annotations

import csv
import datetime
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPORTS_DIR = Path(__file__).resolve().parent / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)


def _apply_sheet_styling(ws: Any, is_newly_activated_sheet: bool = False) -> None:
    """Apply professional colors, headers, borders, and auto-column widths."""
    header_fill = PatternFill(
        start_color="1E4620" if is_newly_activated_sheet else "1F2428",
        end_color="1E4620" if is_newly_activated_sheet else "1F2428",
        fill_type="solid",
    )
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    thin_border = Border(
        left=Side(style="thin", color="E1E4E8"),
        right=Side(style="thin", color="E1E4E8"),
        top=Side(style="thin", color="E1E4E8"),
        bottom=Side(style="thin", color="E1E4E8"),
    )

    active_fill = PatternFill(start_color="E6FFEC", end_color="E6FFEC", fill_type="solid")
    active_font = Font(name="Segoe UI", size=10, bold=True, color="1A7F37")

    inactive_fill = PatternFill(start_color="FFF8C5", end_color="FFF8C5", fill_type="solid")
    inactive_font = Font(name="Segoe UI", size=10, bold=True, color="9A6700")

    pending_fill = PatternFill(start_color="F2EAFA", end_color="F2EAFA", fill_type="solid")
    pending_font = Font(name="Segoe UI", size=10, bold=True, color="6B21A8")

    err_fill = PatternFill(start_color="FFEBE9", end_color="FFEBE9", fill_type="solid")
    err_font = Font(name="Segoe UI", size=10, bold=True, color="CF222E")

    # Format header row
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Format data rows
    headers = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
    active_col_idx = headers.index("Active?") + 1 if "Active?" in headers else -1

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        is_active_val = ws.cell(row=row_idx, column=active_col_idx).value if active_col_idx != -1 else ""

        for col_idx in range(1, ws.max_column + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            c.font = Font(name="Segoe UI", size=10)
            c.alignment = Alignment(vertical="center")

            # Center-align specific columns
            header_name = headers[col_idx - 1]
            if header_name in ["Active?", "Serial / IMEI", "Group", "Last Checked", "Activated At", "Check Status"]:
                c.alignment = Alignment(horizontal="center", vertical="center")

            # Highlight Active column
            if col_idx == active_col_idx:
                act_str = str(is_active_val).strip()
                if act_str == "Yes":
                    c.fill = active_fill
                    c.font = active_font
                elif act_str == "No":
                    c.fill = inactive_fill
                    c.font = inactive_font
                elif act_str in ["Pending", "Pending Sync", "Syncing"]:
                    c.fill = pending_fill
                    c.font = pending_font
                elif act_str in ["Error", "Captcha Failed", "Unknown", "No Response"]:
                    c.fill = err_fill
                    c.font = err_font

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_full_run_sheet(devices: list[dict[str, Any]], custom_tag: str = "") -> tuple[Path, Path]:
    """
    Generate Sheet #1: Full Run Report with all processed serials.
    Returns (excel_path, csv_path).
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    tag_part = f"_{custom_tag.strip()}" if custom_tag.strip() else ""
    filename_base = f"oppo_full_run{tag_part}_{timestamp}"

    excel_path = REPORTS_DIR / f"{filename_base}.xlsx"
    csv_path = REPORTS_DIR / f"{filename_base}.csv"

    columns = [
        ("serial", "Serial / IMEI"),
        ("is_active", "Active?"),
        ("activation_date", "Activation Date"),
        ("activation_comment", "Activation Comment"),
        ("model_name", "Model Name"),
        ("specs", "RAM / Storage"),
        ("color", "Color"),
        ("certified_model", "Certified Model"),
        ("region", "Region"),
        ("group_name", "Group"),
        ("warranty_end", "Warranty End"),
        ("status", "Check Status"),
        ("last_checked_at", "Last Checked"),
        ("activated_at", "Activated At"),
    ]

    # 1. Create Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Full Run Report"

    # Header
    ws.append([col_title for _, col_title in columns])

    # Rows
    for dev in devices:
        row_vals = [dev.get(col_key, "") for col_key, _ in columns]
        ws.append(row_vals)

    _apply_sheet_styling(ws, is_newly_activated_sheet=False)
    wb.save(excel_path)

    # 2. Create Companion CSV
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([col_title for _, col_title in columns])
        for dev in devices:
            writer.writerow([dev.get(col_key, "") for col_key, _ in columns])

    return excel_path, csv_path


def generate_newly_activated_sheet(
    newly_activated_devices: list[dict[str, Any]],
    custom_tag: str = "",
) -> tuple[Path, Path]:
    """
    Generate Sheet #2: Newly Activated Devices Report.
    Only contains devices detected as active in this batch.
    Returns (excel_path, csv_path).
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    tag_part = f"_{custom_tag.strip()}" if custom_tag.strip() else ""
    filename_base = f"oppo_newly_activated{tag_part}_{timestamp}"

    excel_path = REPORTS_DIR / f"{filename_base}.xlsx"
    csv_path = REPORTS_DIR / f"{filename_base}.csv"

    columns = [
        ("serial", "Serial / IMEI"),
        ("is_active", "Active?"),
        ("activation_date", "Activation Date"),
        ("activation_comment", "Activation Comment"),
        ("model_name", "Model Name"),
        ("specs", "RAM / Storage"),
        ("color", "Color"),
        ("group_name", "Group"),
        ("activated_at", "Detected Active At"),
        ("status", "Check Status"),
    ]

    # 1. Create Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Newly Activated"

    # Header
    ws.append([col_title for _, col_title in columns])

    if newly_activated_devices:
        for dev in newly_activated_devices:
            row_vals = [dev.get(col_key, "") for col_key, _ in columns]
            ws.append(row_vals)
    else:
        # Informative empty row
        empty_row = ["No newly activated devices detected in this run", "—", "—", "—", "—", "—", "—", "—", "—", "—"]
        ws.append(empty_row)

    _apply_sheet_styling(ws, is_newly_activated_sheet=True)
    wb.save(excel_path)

    # 2. Create Companion CSV
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([col_title for _, col_title in columns])
        if newly_activated_devices:
            for dev in newly_activated_devices:
                writer.writerow([dev.get(col_key, "") for col_key, _ in columns])
        else:
            writer.writerow(["No newly activated devices detected in this run"] + ["—"] * 9)

    return excel_path, csv_path


def generate_filtered_devices_sheet(
    devices: list[dict[str, Any]],
    custom_tag: str = "filtered",
    sheet_title: str = "Filtered Devices",
) -> tuple[Path, Path]:
    """
    Generate an Excel sheet for filtered devices / selected category.
    Returns (excel_path, csv_path).
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_tag = re.sub(r"[^\w\-]", "_", custom_tag.strip()) or "filtered"
    filename_base = f"oppo_devices_{safe_tag}_{timestamp}"

    excel_path = REPORTS_DIR / f"{filename_base}.xlsx"
    csv_path = REPORTS_DIR / f"{filename_base}.csv"

    columns = [
        ("serial", "Serial / IMEI"),
        ("is_active", "Active?"),
        ("activation_date", "Activation Date"),
        ("activation_comment", "Activation Comment"),
        ("model_name", "Model Name"),
        ("specs", "RAM / Storage"),
        ("color", "Color"),
        ("certified_model", "Certified Model"),
        ("region", "Region"),
        ("group_name", "Group"),
        ("warranty_end", "Warranty End"),
        ("status", "Check Status"),
        ("last_checked_at", "Last Checked"),
        ("activated_at", "Activated At"),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    clean_title = re.sub(r"[\:\\/\?\*\[\]]", "_", sheet_title)[:31].strip() or "Filtered Devices"
    ws.title = clean_title

    ws.append([col_title for _, col_title in columns])

    if devices:
        for dev in devices:
            ws.append([dev.get(col_key, "") for col_key, _ in columns])
    else:
        ws.append(["No devices matched the selected category/filter"] + ["—"] * 13)

    _apply_sheet_styling(ws, is_newly_activated_sheet=False)
    wb.save(excel_path)

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([col_title for _, col_title in columns])
        if devices:
            for dev in devices:
                writer.writerow([dev.get(col_key, "") for col_key, _ in columns])
        else:
            writer.writerow(["No devices matched the selected category/filter"] + ["—"] * 13)

    return excel_path, csv_path


def generate_categorized_workbook(
    categorized_devices: dict[str, list[dict[str, Any]]],
    category_label: str = "Model",
) -> tuple[Path, Path]:
    """
    Generate a multi-sheet Excel workbook with a Summary sheet and separate tabs for each category.
    Returns (excel_path, csv_path).
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_label = re.sub(r"[^\w\-]", "_", category_label.strip()) or "category"
    filename_base = f"oppo_categorized_by_{safe_label.lower()}_{timestamp}"

    excel_path = REPORTS_DIR / f"{filename_base}.xlsx"
    csv_path = REPORTS_DIR / f"{filename_base}.csv"

    columns = [
        ("serial", "Serial / IMEI"),
        ("is_active", "Active?"),
        ("activation_date", "Activation Date"),
        ("activation_comment", "Activation Comment"),
        ("model_name", "Model Name"),
        ("specs", "RAM / Storage"),
        ("color", "Color"),
        ("certified_model", "Certified Model"),
        ("region", "Region"),
        ("group_name", "Group"),
        ("warranty_end", "Warranty End"),
        ("status", "Check Status"),
        ("last_checked_at", "Last Checked"),
        ("activated_at", "Activated At"),
    ]

    wb = openpyxl.Workbook()
    # 1. Summary Sheet
    ws_summary = wb.active
    ws_summary.title = "Summary & Overview"

    summary_headers = [
        f"Category ({category_label})",
        "Total Devices",
        "Active",
        "Inactive",
        "Pending Sync",
        "Not Run",
        "Errors / Other",
        "Active Rate",
    ]
    ws_summary.append(summary_headers)

    summary_rows = []
    tot_total = 0
    tot_active = 0
    tot_inactive = 0
    tot_pending = 0
    tot_not_run = 0
    tot_errors = 0

    for cat_name, devs in categorized_devices.items():
        total = len(devs)
        active = sum(1 for d in devs if str(d.get("is_active")).strip() == "Yes")
        inactive = sum(1 for d in devs if str(d.get("is_active")).strip() == "No")
        pending = sum(
            1 for d in devs if str(d.get("is_active")).strip() in ["Pending", "Pending Sync"]
            or (str(d.get("status")).strip() == "Pending" and str(d.get("is_active")).strip() != "Not Run")
        )
        not_run = sum(
            1 for d in devs if str(d.get("is_active")).strip() == "Not Run"
            or str(d.get("status")).strip() == "Not Run"
            or (not d.get("last_checked_at") and str(d.get("is_active")).strip() not in ["Yes", "No", "Pending"])
        )
        errors = sum(
            1 for d in devs if str(d.get("is_active")).strip() in ["Error", "Unknown"]
            or str(d.get("status")).strip() in ["Error", "Captcha Failed", "No Response"]
        )
        rate_str = f"{(active / total * 100):.1f}%" if total > 0 else "0.0%"

        tot_total += total
        tot_active += active
        tot_inactive += inactive
        tot_pending += pending
        tot_not_run += not_run
        tot_errors += errors

        row = [cat_name, total, active, inactive, pending, not_run, errors, rate_str]
        summary_rows.append(row)
        ws_summary.append(row)

    # Append overall total row
    tot_rate_str = f"{(tot_active / tot_total * 100):.1f}%" if tot_total > 0 else "0.0%"
    total_row = ["TOTAL / ALL CATEGORIES", tot_total, tot_active, tot_inactive, tot_pending, tot_not_run, tot_errors, tot_rate_str]
    ws_summary.append(total_row)

    # Style Summary Sheet
    header_fill = PatternFill(start_color="0F3C28", end_color="0F3C28", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    total_fill = PatternFill(start_color="E6FFEC", end_color="E6FFEC", fill_type="solid")
    total_font = Font(name="Segoe UI", size=11, bold=True, color="0F3C28")
    thin_border = Border(
        left=Side(style="thin", color="E1E4E8"),
        right=Side(style="thin", color="E1E4E8"),
        top=Side(style="thin", color="E1E4E8"),
        bottom=Side(style="thin", color="E1E4E8"),
    )
    double_bottom_border = Border(
        left=Side(style="thin", color="E1E4E8"),
        right=Side(style="thin", color="E1E4E8"),
        top=Side(style="thin", color="E1E4E8"),
        bottom=Side(style="double", color="1F2428"),
    )

    for col_idx in range(1, len(summary_headers) + 1):
        c = ws_summary.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws_summary.row_dimensions[1].height = 28

    for row_idx in range(2, ws_summary.max_row):
        ws_summary.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(summary_headers) + 1):
            c = ws_summary.cell(row=row_idx, column=col_idx)
            c.border = thin_border
            c.font = Font(name="Segoe UI", size=10)
            if col_idx == 1:
                c.alignment = Alignment(horizontal="left", vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Format total row
    last_row_idx = ws_summary.max_row
    ws_summary.row_dimensions[last_row_idx].height = 24
    for col_idx in range(1, len(summary_headers) + 1):
        c = ws_summary.cell(row=last_row_idx, column=col_idx)
        c.fill = total_fill
        c.font = total_font
        c.border = double_bottom_border
        if col_idx == 1:
            c.alignment = Alignment(horizontal="left", vertical="center")
        else:
            c.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws_summary.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # 2. Individual Category Sheets
    used_titles: list[str] = ["Summary & Overview"]
    for cat_name, devs in categorized_devices.items():
        base_title = re.sub(r"[\:\\/\?\*\[\]]", "_", cat_name).strip()[:28] or "Category"
        sheet_title = base_title
        counter = 1
        while sheet_title.lower() in [t.lower() for t in used_titles]:
            counter += 1
            sheet_title = f"{base_title[:25]}_{counter}"[:31]
        used_titles.append(sheet_title)

        ws_cat = wb.create_sheet(title=sheet_title)
        ws_cat.append([col_title for _, col_title in columns])
        for dev in devs:
            ws_cat.append([dev.get(col_key, "") for col_key, _ in columns])
        _apply_sheet_styling(ws_cat, is_newly_activated_sheet=False)

    wb.save(excel_path)

    # Companion summary CSV
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(summary_headers)
        for r in summary_rows:
            writer.writerow(r)
        writer.writerow(total_row)

    return excel_path, csv_path


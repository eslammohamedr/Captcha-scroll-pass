"""OPPO Serial & IMEI Verification GUI Application (Dark Mode Edition).

A professional, high-DPI Windows Desktop application for bulk-checking OPPO devices
with automated Dingxiang slider CAPTCHA solving, live progress tracking,
interactive results table, and formatted Excel reporting.
"""

from __future__ import annotations

import ctypes
import datetime
import os
import queue
import subprocess
import sys
import threading
import time
from pathlib import Path
from typing import Any

# Determine application base path (supports PyInstaller frozen executable)
if getattr(sys, "frozen", False):
    APP_DIR = Path(sys._MEIPASS)
    RUN_DIR = Path(sys.executable).parent
else:
    APP_DIR = Path(__file__).resolve().parent
    RUN_DIR = APP_DIR

# Enable High-DPI awareness on Windows before importing Tkinter
if sys.platform == "win32":
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText

# Core automation imports
try:
    from playwright.sync_api import sync_playwright
    HAS_PLAYWRIGHT = True
except ImportError:
    HAS_PLAYWRIGHT = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Import checker logic from check_oppo_cv
try:
    from check_oppo_cv import check_oppo_serial, read_identifiers, save_to_excel
    HAS_CHECKER = True
except Exception as e:
    HAS_CHECKER = False
    CHECKER_IMPORT_ERROR = str(e)


# ==============================================================================
# Premium Dark Mode Color Palette (GitHub / VS Code Dark Style)
# ==============================================================================
DARK_BG_MAIN = "#0D1117"         # Root window deep background
DARK_BG_SURFACE = "#161B22"      # Cards, panels, sidebars
DARK_BG_SURFACE_ALT = "#21262D"  # Inner containers, table headers
DARK_BORDER = "#30363D"          # Subtle dividers
DARK_INPUT_BG = "#0D1117"        # Text entries and spinboxes
DARK_INPUT_FG = "#F0F6FC"        # Bright text inside inputs

DARK_TEXT_PRIMARY = "#F0F6FC"    # Primary white text
DARK_TEXT_MUTED = "#8B949E"      # Secondary gray labels
DARK_TEXT_DIM = "#6E7681"        # Disabled / subtle hints

DARK_ACCENT_BLUE = "#58A6FF"     # Active highlights & info
DARK_ACCENT_CYAN = "#388BFD"
DARK_BTN_GREEN = "#238636"       # Start button (Vibrant emerald)
DARK_BTN_GREEN_HOVER = "#2EA043"
DARK_BTN_RED = "#DA3633"         # Stop button (Crimson)
DARK_BTN_RED_HOVER = "#F85149"
DARK_BTN_NEUTRAL = "#21262D"     # Secondary action buttons
DARK_BTN_NEUTRAL_HOVER = "#30363D"

# Table row badge colors
DARK_TAG_YES_BG = "#122E1A"
DARK_TAG_YES_FG = "#7EE787"
DARK_TAG_NO_BG = "#38290B"
DARK_TAG_NO_FG = "#F2CC60"
DARK_TAG_PENDING_BG = "#132D48"
DARK_TAG_PENDING_FG = "#79C0FF"
DARK_TAG_FAIL_BG = "#3B181A"
DARK_TAG_FAIL_FG = "#FFA198"


class OppoCheckerGUI:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("OPPO Device Verification Portal — Closed-Loop CAPTCHA Solver")
        self.root.geometry("1220x840")
        self.root.minsize(1020, 700)

        # Apply dark background to root
        self.root.configure(bg=DARK_BG_MAIN)

        # Threading & queues
        self.msg_queue: queue.Queue[tuple[str, Any]] = queue.Queue()
        self.worker_thread: threading.Thread | None = None
        self.stop_requested = threading.Event()
        self.is_running = False

        # Current loaded serials
        self.loaded_serials: list[str] = []
        self.active_output_file: Path = (RUN_DIR / "oppo_results.xlsx").resolve()

        # Stats counters
        self.stat_total = tk.IntVar(value=0)
        self.stat_processed = tk.IntVar(value=0)
        self.stat_activated = tk.IntVar(value=0)
        self.stat_not_activated = tk.IntVar(value=0)
        self.stat_pending = tk.IntVar(value=0)
        self.stat_failed = tk.IntVar(value=0)

        # Settings variables (DEFAULT HEADFUL=TRUE so Chrome is visible by default!)
        self.var_headful = tk.BooleanVar(value=True)
        self.var_skip_verified = tk.BooleanVar(value=True)
        self.var_delay = tk.DoubleVar(value=2.0)
        self.var_limit = tk.StringVar(value="")
        self.var_offset = tk.DoubleVar(value=0.0)
        self.var_output_path = tk.StringVar(value=str(self.active_output_file))

        # Setup custom dark styles
        self._setup_dark_styles()

        # Build UI layout
        self._build_ui()

        # Start periodic queue polling
        self.root.after(80, self._poll_queue)

        # Log system status on launch
        self._initial_system_check()

    def _setup_dark_styles(self) -> None:
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass

        # Global base styling in dark mode
        style.configure(".", font=("Segoe UI", 9), background=DARK_BG_MAIN, foreground=DARK_TEXT_PRIMARY)
        style.configure("TFrame", background=DARK_BG_MAIN)
        style.configure("Surface.TFrame", background=DARK_BG_SURFACE)
        style.configure("TLabel", background=DARK_BG_MAIN, foreground=DARK_TEXT_PRIMARY)
        style.configure("Surface.TLabel", background=DARK_BG_SURFACE, foreground=DARK_TEXT_PRIMARY)

        # Headers
        style.configure("Title.TLabel", font=("Segoe UI", 13, "bold"), foreground=DARK_TEXT_PRIMARY)
        style.configure("Subtitle.TLabel", font=("Segoe UI", 9), foreground=DARK_TEXT_MUTED)
        style.configure("Section.TLabel", font=("Segoe UI", 10, "bold"), foreground=DARK_TEXT_PRIMARY)

        # Stat cards styling
        style.configure("StatNumber.TLabel", font=("Segoe UI", 17, "bold"), foreground=DARK_ACCENT_BLUE, background=DARK_BG_SURFACE)
        style.configure("StatNumberSuccess.TLabel", font=("Segoe UI", 17, "bold"), foreground="#7EE787", background=DARK_BG_SURFACE)
        style.configure("StatNumberWarning.TLabel", font=("Segoe UI", 17, "bold"), foreground="#F2CC60", background=DARK_BG_SURFACE)
        style.configure("StatNumberInfo.TLabel", font=("Segoe UI", 17, "bold"), foreground="#79C0FF", background=DARK_BG_SURFACE)
        style.configure("StatNumberDanger.TLabel", font=("Segoe UI", 17, "bold"), foreground="#FFA198", background=DARK_BG_SURFACE)
        style.configure("StatLabel.TLabel", font=("Segoe UI", 8), foreground=DARK_TEXT_MUTED, background=DARK_BG_SURFACE)

        # Notebook tabs (Dark Mode)
        style.configure("TNotebook", background=DARK_BG_MAIN, borderwidth=0)
        style.configure(
            "TNotebook.Tab",
            font=("Segoe UI", 9, "bold"),
            padding=[14, 6],
            background=DARK_BG_SURFACE_ALT,
            foreground=DARK_TEXT_MUTED,
            borderwidth=0,
        )
        style.map(
            "TNotebook.Tab",
            background=[("selected", DARK_BG_SURFACE)],
            foreground=[("selected", DARK_ACCENT_BLUE)],
        )

        # Treeview table in Dark Mode
        style.configure(
            "Treeview",
            font=("Segoe UI", 9),
            rowheight=28,
            background=DARK_BG_SURFACE,
            fieldbackground=DARK_BG_SURFACE,
            foreground=DARK_TEXT_PRIMARY,
            borderwidth=0,
        )
        style.configure(
            "Treeview.Heading",
            font=("Segoe UI", 9, "bold"),
            background=DARK_BG_SURFACE_ALT,
            foreground=DARK_TEXT_PRIMARY,
            relief="flat",
        )
        style.map(
            "Treeview",
            background=[("selected", "#1F4068")],
            foreground=[("selected", "#FFFFFF")],
        )

        # Progress bar
        style.configure(
            "Accent.Horizontal.TProgressbar",
            troughcolor=DARK_BG_SURFACE_ALT,
            background=DARK_ACCENT_BLUE,
            thickness=12,
            borderwidth=0,
        )

    def _build_ui(self) -> None:
        # Top Header Banner
        banner_frame = tk.Frame(self.root, bg=DARK_BG_SURFACE, height=68, bd=0, highlightthickness=1, highlightbackground=DARK_BORDER)
        banner_frame.pack(side=tk.TOP, fill=tk.X, padx=0, pady=0)
        banner_frame.pack_propagate(False)

        banner_inner = tk.Frame(banner_frame, bg=DARK_BG_SURFACE)
        banner_inner.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)

        # Logo / Title box
        title_box = tk.Frame(banner_inner, bg=DARK_BG_SURFACE)
        title_box.pack(side=tk.LEFT, anchor=tk.W)

        lbl_title = tk.Label(
            title_box,
            text="OPPO Device Verification Portal",
            font=("Segoe UI", 13, "bold"),
            bg=DARK_BG_SURFACE,
            fg=DARK_TEXT_PRIMARY,
        )
        lbl_title.pack(side=tk.LEFT, anchor=tk.W)

        lbl_badge = tk.Label(
            title_box,
            text=" DARK EDITION • v2.0 ",
            font=("Segoe UI", 8, "bold"),
            bg="#238636",
            fg="#FFFFFF",
            padx=6,
            pady=1,
        )
        lbl_badge.pack(side=tk.LEFT, padx=(10, 0))

        lbl_sub = tk.Label(
            banner_inner,
            text="Closed-Loop OpenCV Slider CAPTCHA Solver",
            font=("Segoe UI", 9),
            bg=DARK_BG_SURFACE,
            fg=DARK_TEXT_MUTED,
        )
        lbl_sub.pack(side=tk.LEFT, anchor=tk.W, padx=(8, 0), pady=(2, 0))

        # Top right quick action buttons
        btn_box = tk.Frame(banner_inner, bg=DARK_BG_SURFACE)
        btn_box.pack(side=tk.RIGHT)

        btn_open_excel = tk.Button(
            btn_box,
            text="📊 Open Excel File",
            font=("Segoe UI", 9, "bold"),
            bg="#122E1A",
            fg="#7EE787",
            activebackground="#238636",
            activeforeground="#FFFFFF",
            relief=tk.FLAT,
            padx=12,
            pady=5,
            cursor="hand2",
            command=self.open_excel_file,
        )
        btn_open_excel.pack(side=tk.LEFT, padx=6)

        btn_setup = tk.Button(
            btn_box,
            text="⚙ Setup Browser",
            font=("Segoe UI", 9),
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            activebackground=DARK_BTN_NEUTRAL_HOVER,
            activeforeground=DARK_TEXT_PRIMARY,
            relief=tk.FLAT,
            padx=10,
            pady=5,
            cursor="hand2",
            command=self.setup_playwright_browsers,
        )
        btn_setup.pack(side=tk.LEFT, padx=4)

        # Stat cards row
        stats_frame = tk.Frame(self.root, bg=DARK_BG_MAIN)
        stats_frame.pack(side=tk.TOP, fill=tk.X, padx=16, pady=(10, 6))

        self._create_stat_card(stats_frame, "Total Loaded", self.stat_total, "StatNumber.TLabel", 0)
        self._create_stat_card(stats_frame, "Processed", self.stat_processed, "StatNumberInfo.TLabel", 1)
        self._create_stat_card(stats_frame, "Activated (Yes)", self.stat_activated, "StatNumberSuccess.TLabel", 2)
        self._create_stat_card(stats_frame, "Not Activated (No)", self.stat_not_activated, "StatNumberWarning.TLabel", 3)
        self._create_stat_card(stats_frame, "Pending Sync", self.stat_pending, "StatNumberInfo.TLabel", 4)
        self._create_stat_card(stats_frame, "Errors / Failed", self.stat_failed, "StatNumberDanger.TLabel", 5)

        # Split pane (Left: Input & Controls, Right: Table & Logs)
        main_split = tk.PanedWindow(self.root, orient=tk.HORIZONTAL, bg=DARK_BORDER, sashwidth=4, bd=0)
        main_split.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=16, pady=(6, 14))

        # Left panel (Input & Settings)
        left_container = tk.Frame(main_split, bg=DARK_BG_SURFACE, bd=0, highlightthickness=1, highlightbackground=DARK_BORDER, width=400)
        left_container.pack_propagate(False)
        main_split.add(left_container, minsize=350)

        self._build_left_panel(left_container)

        # Right panel (Live Results & Activity Logs)
        right_container = tk.Frame(main_split, bg=DARK_BG_SURFACE, bd=0, highlightthickness=1, highlightbackground=DARK_BORDER)
        main_split.add(right_container, minsize=520)

        self._build_right_panel(right_container)

    def _create_stat_card(self, parent: tk.Widget, title: str, variable: tk.IntVar, style_name: str, col: int) -> None:
        card = tk.Frame(parent, bg=DARK_BG_SURFACE, bd=0, highlightthickness=1, highlightbackground=DARK_BORDER)
        card.grid(row=0, column=col, sticky="nsew", padx=3, pady=2)
        parent.grid_columnconfigure(col, weight=1)

        inner = tk.Frame(card, bg=DARK_BG_SURFACE)
        inner.pack(fill=tk.BOTH, expand=True, padx=12, pady=8)

        num_lbl = ttk.Label(inner, textvariable=variable, style=style_name)
        num_lbl.pack(anchor=tk.CENTER)

        lbl = ttk.Label(inner, text=title, style="StatLabel.TLabel")
        lbl.pack(anchor=tk.CENTER, pady=(2, 0))

    def _build_left_panel(self, parent: tk.Widget) -> None:
        canvas = tk.Canvas(parent, bg=DARK_BG_SURFACE, bd=0, highlightthickness=0)
        scrollbar = ttk.Scrollbar(parent, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas, bg=DARK_BG_SURFACE)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all")),
        )
        canvas_window = canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")

        def _on_canvas_resize(event):
            canvas.itemconfig(canvas_window, width=event.width)

        canvas.bind("<Configure>", _on_canvas_resize)
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        pad_x = 14

        # ======================================================================
        # SECTION 1: Serials Input
        # ======================================================================
        lbl_sec1 = tk.Label(scrollable_frame, text="1. Select Serials / IMEIs to Check", font=("Segoe UI", 10, "bold"), bg=DARK_BG_SURFACE, fg=DARK_TEXT_PRIMARY)
        lbl_sec1.pack(anchor=tk.W, padx=pad_x, pady=(12, 4))

        input_nb = ttk.Notebook(scrollable_frame)
        input_nb.pack(fill=tk.X, padx=pad_x, pady=2)

        # Tab A: File Upload
        tab_file = tk.Frame(input_nb, bg=DARK_BG_SURFACE, padx=10, pady=10)
        input_nb.add(tab_file, text="📁 File (.txt, .xlsx, .csv)")

        self.lbl_selected_file = tk.Label(
            tab_file,
            text="Current: serials.txt (Auto-detected)",
            font=("Segoe UI", 8),
            bg=DARK_BG_SURFACE,
            fg=DARK_TEXT_MUTED,
            wraplength=310,
            justify=tk.LEFT,
        )
        self.lbl_selected_file.pack(anchor=tk.W, pady=(0, 6))

        file_btn_row = tk.Frame(tab_file, bg=DARK_BG_SURFACE)
        file_btn_row.pack(fill=tk.X)

        btn_browse_serials = tk.Button(
            file_btn_row,
            text="Choose File...",
            font=("Segoe UI", 9),
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            activebackground=DARK_BTN_NEUTRAL_HOVER,
            relief=tk.FLAT,
            padx=10,
            pady=3,
            cursor="hand2",
            command=self.browse_serials_file,
        )
        btn_browse_serials.pack(side=tk.LEFT)

        btn_reload_serials = tk.Button(
            file_btn_row,
            text="🔄 Reload",
            font=("Segoe UI", 9),
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            activebackground=DARK_BTN_NEUTRAL_HOVER,
            relief=tk.FLAT,
            padx=10,
            pady=3,
            cursor="hand2",
            command=self.load_serials_from_default,
        )
        btn_reload_serials.pack(side=tk.LEFT, padx=6)

        # Tab B: Direct Paste
        tab_paste = tk.Frame(input_nb, bg=DARK_BG_SURFACE, padx=10, pady=10)
        input_nb.add(tab_paste, text="📝 Paste Serials")

        lbl_paste_hint = tk.Label(tab_paste, text="Paste one serial or IMEI per line:", font=("Segoe UI", 8), bg=DARK_BG_SURFACE, fg=DARK_TEXT_MUTED)
        lbl_paste_hint.pack(anchor=tk.W, pady=(0, 4))

        self.txt_paste = ScrolledText(
            tab_paste,
            height=6,
            font=("Consolas", 9),
            bg=DARK_INPUT_BG,
            fg=DARK_INPUT_FG,
            insertbackground="#FFFFFF",
            relief=tk.SOLID,
            bd=1,
            highlightbackground=DARK_BORDER,
        )
        self.txt_paste.pack(fill=tk.BOTH, expand=True)
        self.txt_paste.bind("<KeyRelease>", self._on_paste_text_changed)

        # Serials count status
        self.lbl_serials_summary = tk.Label(
            scrollable_frame,
            text="0 serials loaded and ready",
            font=("Segoe UI", 9, "bold"),
            bg=DARK_BG_SURFACE,
            fg=DARK_ACCENT_BLUE,
        )
        self.lbl_serials_summary.pack(anchor=tk.W, padx=pad_x, pady=(6, 8))

        # ======================================================================
        # SECTION 2: Browser & Check Settings (SHOW CHROME PROMINENTLY)
        # ======================================================================
        lbl_sec2 = tk.Label(scrollable_frame, text="2. Browser & Verification Settings", font=("Segoe UI", 10, "bold"), bg=DARK_BG_SURFACE, fg=DARK_TEXT_PRIMARY)
        lbl_sec2.pack(anchor=tk.W, padx=pad_x, pady=(8, 4))

        settings_box = tk.Frame(scrollable_frame, bg=DARK_BG_SURFACE_ALT, bd=0, highlightthickness=1, highlightbackground=DARK_BORDER)
        settings_box.pack(fill=tk.X, padx=pad_x, pady=2)

        # --- FEATURE: SHOW CHROME PROMINENT TOGGLE ---
        chrome_highlight_box = tk.Frame(settings_box, bg="#132D48", bd=0, highlightthickness=1, highlightbackground="#1F4B78")
        chrome_highlight_box.pack(fill=tk.X, padx=8, pady=(8, 6))

        chk_headful = tk.Checkbutton(
            chrome_highlight_box,
            text="👁️  Show Chrome Browser (Watch Live)",
            variable=self.var_headful,
            bg="#132D48",
            fg="#79C0FF",
            activebackground="#132D48",
            activeforeground="#79C0FF",
            selectcolor=DARK_BG_MAIN,
            font=("Segoe UI", 9, "bold"),
            cursor="hand2",
        )
        chk_headful.pack(anchor=tk.W, padx=8, pady=(6, 2))

        lbl_headful_desc = tk.Label(
            chrome_highlight_box,
            text="Keeps Chrome visible on screen so you can observe the portal and CAPTCHA solving live.",
            font=("Segoe UI", 8),
            bg="#132D48",
            fg="#A5D6FF",
            wraplength=310,
            justify=tk.LEFT,
        )
        lbl_headful_desc.pack(anchor=tk.W, padx=8, pady=(0, 6))

        # Skip verified checkbox
        chk_skip = tk.Checkbutton(
            settings_box,
            text="Skip already verified records in Excel",
            variable=self.var_skip_verified,
            bg=DARK_BG_SURFACE_ALT,
            fg=DARK_TEXT_PRIMARY,
            activebackground=DARK_BG_SURFACE_ALT,
            activeforeground=DARK_TEXT_PRIMARY,
            selectcolor=DARK_BG_MAIN,
            font=("Segoe UI", 9),
            cursor="hand2",
        )
        chk_skip.pack(anchor=tk.W, padx=8, pady=3)

        # Delay slider / spinbox
        delay_row = tk.Frame(settings_box, bg=DARK_BG_SURFACE_ALT)
        delay_row.pack(fill=tk.X, padx=8, pady=4)
        tk.Label(delay_row, text="Delay between devices (s):", font=("Segoe UI", 9), bg=DARK_BG_SURFACE_ALT, fg=DARK_TEXT_PRIMARY).pack(side=tk.LEFT)
        spn_delay = tk.Spinbox(
            delay_row,
            from_=0.5,
            to=15.0,
            increment=0.5,
            textvariable=self.var_delay,
            width=5,
            font=("Segoe UI", 9),
            bg=DARK_INPUT_BG,
            fg=DARK_INPUT_FG,
            buttonbackground=DARK_BG_SURFACE,
            relief=tk.FLAT,
        )
        spn_delay.pack(side=tk.RIGHT)

        # Batch limit
        limit_row = tk.Frame(settings_box, bg=DARK_BG_SURFACE_ALT)
        limit_row.pack(fill=tk.X, padx=8, pady=4)
        tk.Label(limit_row, text="Limit checks (e.g. 5, blank=all):", font=("Segoe UI", 9), bg=DARK_BG_SURFACE_ALT, fg=DARK_TEXT_PRIMARY).pack(side=tk.LEFT)
        ent_limit = tk.Entry(
            limit_row,
            textvariable=self.var_limit,
            width=6,
            font=("Segoe UI", 9),
            bg=DARK_INPUT_BG,
            fg=DARK_INPUT_FG,
            insertbackground="#FFFFFF",
            relief=tk.FLAT,
        )
        ent_limit.pack(side=tk.RIGHT)

        # Fine-tune offset
        offset_row = tk.Frame(settings_box, bg=DARK_BG_SURFACE_ALT)
        offset_row.pack(fill=tk.X, padx=8, pady=(4, 8))
        tk.Label(offset_row, text="Slider Offset nudge (px):", font=("Segoe UI", 9), bg=DARK_BG_SURFACE_ALT, fg=DARK_TEXT_PRIMARY).pack(side=tk.LEFT)
        spn_offset = tk.Spinbox(
            offset_row,
            from_=-15.0,
            to=15.0,
            increment=0.5,
            textvariable=self.var_offset,
            width=5,
            font=("Segoe UI", 9),
            bg=DARK_INPUT_BG,
            fg=DARK_INPUT_FG,
            buttonbackground=DARK_BG_SURFACE,
            relief=tk.FLAT,
        )
        spn_offset.pack(side=tk.RIGHT)

        # ======================================================================
        # SECTION 3: Output File
        # ======================================================================
        lbl_sec3 = tk.Label(scrollable_frame, text="3. Target Excel Report File", font=("Segoe UI", 10, "bold"), bg=DARK_BG_SURFACE, fg=DARK_TEXT_PRIMARY)
        lbl_sec3.pack(anchor=tk.W, padx=pad_x, pady=(10, 4))

        out_box = tk.Frame(scrollable_frame, bg=DARK_BG_SURFACE)
        out_box.pack(fill=tk.X, padx=pad_x, pady=2)

        ent_output = tk.Entry(
            out_box,
            textvariable=self.var_output_path,
            font=("Segoe UI", 9),
            bg=DARK_INPUT_BG,
            fg=DARK_INPUT_FG,
            insertbackground="#FFFFFF",
            relief=tk.SOLID,
            bd=1,
            highlightbackground=DARK_BORDER,
        )
        ent_output.pack(side=tk.LEFT, fill=tk.X, expand=True)

        btn_browse_out = tk.Button(
            out_box,
            text="Browse",
            font=("Segoe UI", 8),
            command=self.browse_output_file,
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            relief=tk.FLAT,
            padx=8,
            cursor="hand2",
        )
        btn_browse_out.pack(side=tk.LEFT, padx=(6, 0))

        # ======================================================================
        # SECTION 4: Action Controls
        # ======================================================================
        lbl_sec4 = tk.Label(scrollable_frame, text="4. Execution Controls", font=("Segoe UI", 10, "bold"), bg=DARK_BG_SURFACE, fg=DARK_TEXT_PRIMARY)
        lbl_sec4.pack(anchor=tk.W, padx=pad_x, pady=(12, 4))

        # Big Start Button
        self.btn_start = tk.Button(
            scrollable_frame,
            text="▶   START VERIFICATION",
            font=("Segoe UI", 11, "bold"),
            bg=DARK_BTN_GREEN,
            fg="#FFFFFF",
            activebackground=DARK_BTN_GREEN_HOVER,
            activeforeground="#FFFFFF",
            relief=tk.FLAT,
            pady=9,
            cursor="hand2",
            command=self.start_verification,
        )
        self.btn_start.pack(fill=tk.X, padx=pad_x, pady=4)

        # Stop Button
        self.btn_stop = tk.Button(
            scrollable_frame,
            text="⏹   STOP / CANCEL",
            font=("Segoe UI", 10, "bold"),
            bg=DARK_BTN_RED,
            fg="#FFFFFF",
            activebackground=DARK_BTN_RED_HOVER,
            activeforeground="#FFFFFF",
            relief=tk.FLAT,
            pady=6,
            cursor="hand2",
            state=tk.DISABLED,
            command=self.stop_verification,
        )
        self.btn_stop.pack(fill=tk.X, padx=pad_x, pady=4)

        # Quick tools row
        tools_row = tk.Frame(scrollable_frame, bg=DARK_BG_SURFACE)
        tools_row.pack(fill=tk.X, padx=pad_x, pady=(6, 16))

        btn_clear_table = tk.Button(
            tools_row,
            text="🧹 Clear Table",
            font=("Segoe UI", 8),
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            relief=tk.FLAT,
            cursor="hand2",
            command=self.clear_table,
        )
        btn_clear_table.pack(side=tk.LEFT, expand=True, fill=tk.X, padx=2)

        btn_refresh_counts = tk.Button(
            tools_row,
            text="🔄 Sync Stats",
            font=("Segoe UI", 8),
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            relief=tk.FLAT,
            cursor="hand2",
            command=self.sync_stats_from_excel,
        )
        btn_refresh_counts.pack(side=tk.LEFT, expand=True, fill=tk.X, padx=2)

        # Automatically try to load default serials.txt
        self.load_serials_from_default()

    def _build_right_panel(self, parent: tk.Widget) -> None:
        # Progress bar & current status at top
        progress_box = tk.Frame(parent, bg=DARK_BG_SURFACE, padx=14, pady=10)
        progress_box.pack(side=tk.TOP, fill=tk.X)

        self.lbl_progress_status = tk.Label(
            progress_box,
            text="Ready. Select serials and click 'START VERIFICATION'.",
            font=("Segoe UI", 9, "bold"),
            bg=DARK_BG_SURFACE,
            fg=DARK_TEXT_PRIMARY,
            anchor=tk.W,
        )
        self.lbl_progress_status.pack(fill=tk.X, pady=(0, 4))

        self.progress_bar = ttk.Progressbar(
            progress_box,
            style="Accent.Horizontal.TProgressbar",
            mode="determinate",
            maximum=100,
            value=0,
        )
        self.progress_bar.pack(fill=tk.X)

        # Tabbed view for Results & Console Log
        notebook = ttk.Notebook(parent)
        notebook.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=8, pady=(4, 8))

        # ======================================================================
        # Tab 1: Live Results Table
        # ======================================================================
        tab_results = tk.Frame(notebook, bg=DARK_BG_SURFACE)
        notebook.add(tab_results, text="📋 Live Results Table")

        columns = (
            "sn",
            "active",
            "act_date",
            "act_comment",
            "model",
            "color",
            "status",
            "checked_at",
        )
        self.tree = ttk.Treeview(tab_results, columns=columns, show="headings", selectmode="browse")

        self.tree.heading("sn", text="Serial / IMEI")
        self.tree.heading("active", text="Active?")
        self.tree.heading("act_date", text="Activation Date")
        self.tree.heading("act_comment", text="Activation Comment")
        self.tree.heading("model", text="Model Name")
        self.tree.heading("color", text="Color")
        self.tree.heading("status", text="Status")
        self.tree.heading("checked_at", text="Checked At")

        self.tree.column("sn", width=140, anchor=tk.W)
        self.tree.column("active", width=75, anchor=tk.CENTER)
        self.tree.column("act_date", width=110, anchor=tk.CENTER)
        self.tree.column("act_comment", width=220, anchor=tk.W)
        self.tree.column("model", width=120, anchor=tk.W)
        self.tree.column("color", width=90, anchor=tk.W)
        self.tree.column("status", width=95, anchor=tk.CENTER)
        self.tree.column("checked_at", width=115, anchor=tk.CENTER)

        # Row color tags (Dark Mode palette)
        self.tree.tag_configure("tag_yes", background=DARK_TAG_YES_BG, foreground=DARK_TAG_YES_FG)
        self.tree.tag_configure("tag_no", background=DARK_TAG_NO_BG, foreground=DARK_TAG_NO_FG)
        self.tree.tag_configure("tag_pending", background=DARK_TAG_PENDING_BG, foreground=DARK_TAG_PENDING_FG)
        self.tree.tag_configure("tag_fail", background=DARK_TAG_FAIL_BG, foreground=DARK_TAG_FAIL_FG)
        self.tree.tag_configure("tag_info", background=DARK_BG_SURFACE, foreground=DARK_TEXT_PRIMARY)

        tree_y_scroll = ttk.Scrollbar(tab_results, orient="vertical", command=self.tree.yview)
        tree_x_scroll = ttk.Scrollbar(tab_results, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=tree_y_scroll.set, xscrollcommand=tree_x_scroll.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        tree_y_scroll.grid(row=0, column=1, sticky="ns")
        tree_x_scroll.grid(row=1, column=0, sticky="ew")

        tab_results.grid_rowconfigure(0, weight=1)
        tab_results.grid_columnconfigure(0, weight=1)

        # Double click to view full detail popup
        self.tree.bind("<Double-1>", self._on_tree_row_double_click)

        # ======================================================================
        # Tab 2: Activity Log
        # ======================================================================
        tab_logs = tk.Frame(notebook, bg=DARK_BG_SURFACE)
        notebook.add(tab_logs, text="📜 Activity Log")

        self.txt_log = ScrolledText(
            tab_logs,
            wrap=tk.WORD,
            font=("Consolas", 9),
            bg="#0B0E14",
            fg="#C9D1D9",
            insertbackground="#58A6FF",
            relief=tk.FLAT,
        )
        self.txt_log.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)

        # Log color styles
        self.txt_log.tag_configure("info", foreground="#58A6FF")
        self.txt_log.tag_configure("success", foreground="#7EE787")
        self.txt_log.tag_configure("warning", foreground="#F2CC60")
        self.txt_log.tag_configure("error", foreground="#FFA198")
        self.txt_log.tag_configure("highlight", foreground="#D2A8FF")

        # Bottom log toolbar
        log_bar = tk.Frame(tab_logs, bg=DARK_BG_SURFACE)
        log_bar.pack(fill=tk.X, padx=4, pady=(0, 4))

        btn_clear_log = tk.Button(
            log_bar,
            text="Clear Log",
            font=("Segoe UI", 8),
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            relief=tk.FLAT,
            command=lambda: self.txt_log.delete("1.0", tk.END),
        )
        btn_clear_log.pack(side=tk.RIGHT)

    def log(self, message: str, level: str = "info") -> None:
        """Post a thread-safe message to the log viewer."""
        self.msg_queue.put(("LOG", (level, message)))

    def _initial_system_check(self) -> None:
        """Check prerequisites on start and prompt for automatic installation if missing."""
        self.log("Initializing OPPO Automation Application v2.0 (Dark Edition)...", "info")

        missing: list[str] = []
        if not HAS_PLAYWRIGHT:
            missing.append("playwright")
        if not HAS_OPENPYXL:
            missing.append("openpyxl")
        if not HAS_CHECKER:
            missing.append("opencv-python / check_oppo_cv")

        if missing:
            self.log(f"[Notice] Missing libraries detected: {', '.join(missing)}", "warning")
            if messagebox.askyesno(
                "Install Required Dependencies",
                f"The following required Python libraries are missing:\n\n"
                f"• {chr(10).join(missing)}\n\n"
                "Would you like to install them and download the browser automatically now?",
            ):
                self.install_all_dependencies()
        else:
            self.log("[Ready] All required Python libraries and OpenCV engine loaded successfully.", "success")

        # Initial stats check from existing excel if available
        self.sync_stats_from_excel()

    def install_all_dependencies(self) -> None:
        """Run pip install -r requirements.txt and playwright install in background."""
        self.log("Starting automatic installation of requirements...", "info")

        def _worker():
            try:
                req_path = RUN_DIR / "requirements.txt"
                if not req_path.exists():
                    req_path = APP_DIR / "requirements.txt"

                # 1. pip install -r requirements.txt
                cmd1 = [sys.executable, "-m", "pip", "install", "-r", str(req_path)]
                proc1 = subprocess.Popen(
                    cmd1,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    bufsize=1,
                    creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
                )
                if proc1.stdout:
                    for line in proc1.stdout:
                        l = line.strip()
                        if l:
                            self.msg_queue.put(("LOG", ("info", f"[Pip] {l}")))
                proc1.wait()

                # 2. playwright install chromium
                self._run_playwright_install()

                self.msg_queue.put(("LOG", ("success", "All dependencies installed successfully!")))
                self.msg_queue.put(("POPUP_INFO", ("Installation Complete", "All dependencies and Playwright Chromium have been installed successfully!\nPlease restart the app to apply changes.")))
            except Exception as e:
                self.msg_queue.put(("LOG", ("error", f"Installation failed: {e}")))

        threading.Thread(target=_worker, daemon=True).start()

    def _run_playwright_install(self) -> int:
        """Run playwright install chromium supporting both regular Python and PyInstaller frozen EXE."""
        try:
            # Check if we can use playwright internal driver directly
            from playwright._impl._driver import compute_driver_executable, get_driver_env
            node_exe, cli_js = compute_driver_executable()
            cmd = [str(node_exe), str(cli_js), "install", "chromium"]
            env = get_driver_env()
        except Exception:
            cmd = [sys.executable, "-m", "playwright", "install", "chromium"]
            env = os.environ.copy()

        proc = subprocess.Popen(
            cmd,
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
        )
        if proc.stdout:
            for line in proc.stdout:
                clean_line = line.strip()
                if clean_line:
                    self.msg_queue.put(("LOG", ("info", f"[Browser Setup] {clean_line}")))
        proc.wait()
        return proc.returncode

    def browse_serials_file(self) -> None:
        f = filedialog.askopenfilename(
            title="Select Serials / IMEI Input File",
            filetypes=[("All Supported Files", "*.txt *.xlsx *.csv"), ("Text Files", "*.txt"), ("Excel Files", "*.xlsx"), ("CSV Files", "*.csv")],
        )
        if not f:
            return

        try:
            path = Path(f)
            serials = read_identifiers(path)
            self.loaded_serials = serials
            self.lbl_selected_file.config(text=f"Current: {path.name} ({len(serials)} serials)")
            self.lbl_serials_summary.config(text=f"{len(serials)} serials loaded from {path.name}")
            self.stat_total.set(len(serials))
            self.log(f"Loaded {len(serials)} serial(s) from {path}", "info")
        except Exception as e:
            messagebox.showerror("File Error", f"Could not read input file:\n{e}")
            self.log(f"Failed to read file {f}: {e}", "error")

    def load_serials_from_default(self) -> None:
        default_p = RUN_DIR / "serials.txt"
        if not default_p.exists():
            default_p = APP_DIR / "serials.txt"

        if default_p.exists():
            try:
                serials = read_identifiers(default_p)
                self.loaded_serials = serials
                self.lbl_selected_file.config(text=f"Current: serials.txt ({len(serials)} serials)")
                self.lbl_serials_summary.config(text=f"{len(serials)} serials loaded from serials.txt")
                self.stat_total.set(len(serials))
                self.log(f"Auto-loaded {len(serials)} serials from serials.txt", "info")
            except Exception as e:
                self.log(f"Could not auto-load serials.txt: {e}", "warning")

    def _on_paste_text_changed(self, event: Any = None) -> None:
        content = self.txt_paste.get("1.0", tk.END).strip()
        if not content:
            return
        lines = [line.strip() for line in content.splitlines() if line.strip() and not line.strip().startswith("#")]
        self.loaded_serials = lines
        self.lbl_serials_summary.config(text=f"{len(lines)} serials ready from pasted text")
        self.stat_total.set(len(lines))

    def browse_output_file(self) -> None:
        f = filedialog.asksaveasfilename(
            title="Choose Output Excel File",
            defaultextension=".xlsx",
            initialfile="oppo_results.xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
        )
        if f:
            self.var_output_path.set(f)
            self.active_output_file = Path(f)
            self.log(f"Output Excel set to: {f}", "info")
            self.sync_stats_from_excel()

    def open_excel_file(self) -> None:
        """Open the generated Excel spreadsheet with the default Windows application."""
        p = Path(self.var_output_path.get()).resolve()
        if not p.exists():
            messagebox.showinfo("Excel File Not Found", f"The file '{p.name}' does not exist yet.\nRun verification first to generate it.")
            return

        try:
            os.startfile(str(p))
            self.log(f"Opened Excel report: {p}", "success")
        except Exception as e:
            messagebox.showerror("Open File Error", f"Could not open file in Excel:\n{e}")

    def clear_table(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.log("Results table cleared.", "info")

    def sync_stats_from_excel(self) -> None:
        """Read existing Excel output and update the counters."""
        p = Path(self.var_output_path.get())
        if not p.exists() or not HAS_OPENPYXL:
            return

        try:
            wb = openpyxl.load_workbook(p, data_only=True)
            ws = wb.active
            headers = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []

            act_col = headers.index("Active?") if "Active?" in headers else -1
            stat_col = headers.index("Check Status") if "Check Status" in headers else -1

            yes_count = 0
            no_count = 0
            pending_count = 0
            failed_count = 0
            total_records = 0

            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                total_records += 1
                act_val = str(row[act_col] or "").strip() if act_col != -1 and len(row) > act_col else ""
                stat_val = str(row[stat_col] or "").strip() if stat_col != -1 and len(row) > stat_col else ""

                if act_val == "Yes":
                    yes_count += 1
                elif act_val == "No":
                    no_count += 1
                elif act_val in ["Pending", "Syncing"]:
                    pending_count += 1
                elif stat_val in ["Captcha Failed", "Error", "Not Found"]:
                    failed_count += 1

            self.stat_activated.set(yes_count)
            self.stat_not_activated.set(no_count)
            self.stat_pending.set(pending_count)
            self.stat_failed.set(failed_count)
            self.log(f"Synced stats from {p.name}: {total_records} saved records found.", "info")
        except Exception as e:
            self.log(f"Notice: Could not read {p.name} for stats sync: {e}", "warning")

    def setup_playwright_browsers(self) -> None:
        """One-click tool to install Chromium browser for customers."""
        if messagebox.askyesno("Setup Browser", "This will download and install the required Chromium browser engine (~150MB).\n\nDo you want to proceed?"):
            self.log("Starting Chromium browser setup in the background...", "info")

            def _install_worker():
                code = self._run_playwright_install()
                if code == 0:
                    self.msg_queue.put(("LOG", ("success", "Chromium browser setup completed successfully!")))
                    self.msg_queue.put(("POPUP_INFO", ("Setup Complete", "Chromium browser installed and ready!")))
                else:
                    self.msg_queue.put(("LOG", ("error", f"Browser setup exited with error code {code}")))

            threading.Thread(target=_install_worker, daemon=True).start()

    def start_verification(self) -> None:
        """Initiate background verification thread."""
        if self.is_running:
            return

        paste_text = self.txt_paste.get("1.0", tk.END).strip()
        if paste_text:
            serials = [line.strip() for line in paste_text.splitlines() if line.strip() and not line.strip().startswith("#")]
        else:
            serials = list(self.loaded_serials)

        if not serials:
            messagebox.showwarning("No Serials", "Please choose a file or paste serial numbers before starting.")
            return

        limit_val = self.var_limit.get().strip()
        if limit_val.isdigit():
            lim = int(limit_val)
            serials = serials[:lim]

        self.stop_requested.clear()
        self.is_running = True
        self.btn_start.config(state=tk.DISABLED, bg="#484F58")
        self.btn_stop.config(state=tk.NORMAL)
        self.progress_bar.config(value=0)
        self.stat_processed.set(0)

        params = {
            "serials": serials,
            "headful": self.var_headful.get(),
            "skip_verified": self.var_skip_verified.get(),
            "delay": self.var_delay.get(),
            "offset": self.var_offset.get(),
            "output_path": Path(self.var_output_path.get()),
        }

        self.worker_thread = threading.Thread(target=self._worker_run, args=(params,), daemon=True)
        self.worker_thread.start()

    def stop_verification(self) -> None:
        """Signal background worker to stop gracefully."""
        if not self.is_running:
            return
        self.log("Stopping... Finishing current device operation safely.", "warning")
        self.lbl_progress_status.config(text="Stopping... Please wait for current device to finish.")
        self.stop_requested.set()
        self.btn_stop.config(state=tk.DISABLED)

    def _worker_run(self, params: dict[str, Any]) -> None:
        """Background thread logic running Playwright and saving Excel."""
        serials: list[str] = params["serials"]
        headful: bool = params["headful"]
        skip_verified: bool = params["skip_verified"]
        delay: float = params["delay"]
        offset: float = params["offset"]
        output_path: Path = params["output_path"]

        self.msg_queue.put(("LOG", ("info", f"Started check batch. Target serials: {len(serials)} | Browser: {'VISIBLE (Watch Live)' if headful else 'Background Headless'}")))

        # Check skip verified
        if skip_verified and output_path.exists() and HAS_OPENPYXL:
            try:
                wb = openpyxl.load_workbook(output_path, data_only=True)
                ws = wb.active
                headers = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
                stat_col = headers.index("Check Status") if "Check Status" in headers else -1
                sn_col = headers.index("Serial / IMEI") if "Serial / IMEI" in headers else 0
                found_sns = set()
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row and stat_col != -1 and len(row) > stat_col and row[stat_col] == "Found":
                        found_sns.add(str(row[sn_col]).strip())
                original_len = len(serials)
                serials = [s for s in serials if s not in found_sns]
                skipped_count = original_len - len(serials)
                if skipped_count > 0:
                    self.msg_queue.put(("LOG", ("info", f"Skipping {skipped_count} serials already marked as 'Found' in Excel.")))
            except Exception as e:
                self.msg_queue.put(("LOG", ("warning", f"Could not check existing records: {e}")))

        total = len(serials)
        if total == 0:
            self.msg_queue.put(("LOG", ("success", "All serials are already verified in Excel!")))
            self.msg_queue.put(("FINISHED", {"total": 0, "processed": 0}))
            return

        results_accum: list[dict[str, Any]] = []

        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(
                    headless=not headful,
                    args=["--start-maximized"] if headful else [],
                )
                context = browser.new_context(
                    user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
                    viewport={"width": 1280, "height": 820},
                )
                page = context.new_page()

                for idx, sn in enumerate(serials, start=1):
                    if self.stop_requested.is_set():
                        self.msg_queue.put(("LOG", ("warning", f"Batch stopped by user after {idx-1} checks.")))
                        break

                    pct = int(((idx - 1) / total) * 100)
                    self.msg_queue.put(("PROGRESS", (pct, f"[{idx}/{total}] Checking {sn}...")))

                    def on_status_cb(msg: str):
                        self.msg_queue.put(("PROGRESS_STATUS", f"[{idx}/{total}] {sn}: {msg}"))

                    res = check_oppo_serial(page, sn, on_status=on_status_cb, fine_tune_offset=offset)
                    results_accum.append(res)

                    # Update Excel safely with permission error retry handling
                    self._safe_save_excel(results_accum, output_path)

                    # Send result to UI
                    self.msg_queue.put(("ROW_RESULT", res))

                    log_msg = f"[{idx}/{total}] {sn} -> Active: {res.get('is_active')} | Model: {res.get('model_name') or 'N/A'}"
                    if res.get("is_active") == "Yes":
                        self.msg_queue.put(("LOG", ("success", log_msg)))
                    elif res.get("is_active") == "No":
                        self.msg_queue.put(("LOG", ("warning", log_msg)))
                    else:
                        self.msg_queue.put(("LOG", ("info", log_msg)))

                    # Update progress bar
                    current_pct = int((idx / total) * 100)
                    self.msg_queue.put(("PROGRESS", (current_pct, f"[{idx}/{total}] Completed {sn}")))

                    # Delay between checks if not last item
                    if idx < total and not self.stop_requested.is_set():
                        time.sleep(delay)

                browser.close()
        except Exception as e:
            self.msg_queue.put(("LOG", ("error", f"Automation engine error: {e}")))
            self.msg_queue.put(("POPUP_ERROR", ("Error", f"Automation encountered an error:\n{e}")))

        # Final save
        if results_accum:
            self._safe_save_excel(results_accum, output_path)

        self.msg_queue.put(("FINISHED", {"total": total, "processed": len(results_accum)}))

    def _safe_save_excel(self, results: list[dict[str, Any]], output_path: Path) -> None:
        """Save results to Excel with retry on Windows file lock."""
        for attempt in range(3):
            try:
                save_to_excel(results, output_path, quiet=True)
                return
            except PermissionError:
                if attempt == 0:
                    self.msg_queue.put(("LOG", ("warning", f"[Notice] {output_path.name} is open in Excel! Please close it to allow saving.")))
                time.sleep(1.0)
            except Exception as e:
                self.msg_queue.put(("LOG", ("error", f"Error saving Excel: {e}")))
                break

    def _poll_queue(self) -> None:
        """Process background worker messages and update UI widgets."""
        while not self.msg_queue.empty():
            try:
                msg_type, payload = self.msg_queue.get_nowait()
            except queue.Empty:
                break

            if msg_type == "LOG":
                level, text = payload
                t_str = datetime.datetime.now().strftime("%H:%M:%S")
                self.txt_log.insert(tk.END, f"[{t_str}] {text}\n", level)
                self.txt_log.see(tk.END)

            elif msg_type == "PROGRESS":
                pct, status_text = payload
                self.progress_bar.config(value=pct)
                self.lbl_progress_status.config(text=status_text)

            elif msg_type == "PROGRESS_STATUS":
                self.lbl_progress_status.config(text=payload)

            elif msg_type == "ROW_RESULT":
                res = payload
                sn = res.get("serial", "")
                act = res.get("is_active", "")
                act_date = res.get("activation_date", "")
                act_comment = res.get("activation_comment", "")
                model = res.get("model_name", "")
                color = res.get("color", "")
                status = res.get("status", "")
                t_now = datetime.datetime.now().strftime("%H:%M:%S")

                tag = "tag_info"
                if act == "Yes":
                    tag = "tag_yes"
                    self.stat_activated.set(self.stat_activated.get() + 1)
                elif act == "No":
                    tag = "tag_no"
                    self.stat_not_activated.set(self.stat_not_activated.get() + 1)
                elif act in ["Pending", "Syncing"]:
                    tag = "tag_pending"
                    self.stat_pending.set(self.stat_pending.get() + 1)
                elif status in ["Captcha Failed", "Error", "Not Found"]:
                    tag = "tag_fail"
                    self.stat_failed.set(self.stat_failed.get() + 1)

                self.stat_processed.set(self.stat_processed.get() + 1)

                self.tree.insert(
                    "",
                    0,
                    values=(sn, act, act_date, act_comment, model, color, status, t_now),
                    tags=(tag,),
                )

            elif msg_type == "FINISHED":
                self.is_running = False
                self.btn_start.config(state=tk.NORMAL, bg=DARK_BTN_GREEN)
                self.btn_stop.config(state=tk.DISABLED)
                self.lbl_progress_status.config(text="Verification completed successfully! Click 'Open Excel' to view report.")
                self.progress_bar.config(value=100)
                self.log("Batch finished. Excel spreadsheet updated.", "success")
                messagebox.showinfo("Done", f"Verification completed!\nProcessed {payload.get('processed', 0)} of {payload.get('total', 0)} devices.")

            elif msg_type == "POPUP_INFO":
                title, msg = payload
                messagebox.showinfo(title, msg)

            elif msg_type == "POPUP_ERROR":
                title, msg = payload
                messagebox.showerror(title, msg)

        self.root.after(80, self._poll_queue)

    def _on_tree_row_double_click(self, event: Any) -> None:
        """Display dark detail dialog when user double clicks a table record."""
        selected = self.tree.selection()
        if not selected:
            return
        values = self.tree.item(selected[0], "values")
        if not values:
            return

        detail_win = tk.Toplevel(self.root)
        detail_win.title(f"Device Details - {values[0]}")
        detail_win.geometry("540x440")
        detail_win.minsize(460, 360)
        detail_win.configure(bg=DARK_BG_SURFACE)

        lbl_head = tk.Label(detail_win, text=f"OPPO Device Record: {values[0]}", font=("Segoe UI", 11, "bold"), bg=DARK_BG_SURFACE, fg=DARK_ACCENT_BLUE)
        lbl_head.pack(anchor=tk.W, padx=18, pady=(16, 8))

        fields = [
            ("Serial / IMEI:", values[0]),
            ("Active Status:", values[1]),
            ("Activation Date:", values[2]),
            ("Activation Comment:", values[3]),
            ("Model Name:", values[4]),
            ("Color:", values[5]),
            ("Check Status:", values[6]),
            ("Checked Timestamp:", values[7]),
        ]

        f_frame = tk.Frame(detail_win, bg=DARK_BG_SURFACE)
        f_frame.pack(fill=tk.BOTH, expand=True, padx=18, pady=4)

        for lbl_txt, val_txt in fields:
            row_f = tk.Frame(f_frame, bg=DARK_BG_SURFACE)
            row_f.pack(fill=tk.X, pady=4)

            tk.Label(row_f, text=lbl_txt, width=18, anchor=tk.W, font=("Segoe UI", 9, "bold"), bg=DARK_BG_SURFACE, fg=DARK_TEXT_PRIMARY).pack(side=tk.LEFT)

            val_lbl = tk.Label(row_f, text=str(val_txt), anchor=tk.W, justify=tk.LEFT, wraplength=330, font=("Segoe UI", 9), bg=DARK_BG_SURFACE, fg=DARK_TEXT_MUTED)
            val_lbl.pack(side=tk.LEFT, fill=tk.X, expand=True)

        btn_close = tk.Button(
            detail_win,
            text="Close",
            font=("Segoe UI", 9),
            relief=tk.FLAT,
            bg=DARK_BTN_NEUTRAL,
            fg=DARK_TEXT_PRIMARY,
            command=detail_win.destroy,
            padx=14,
            pady=5,
            cursor="hand2",
        )
        btn_close.pack(side=tk.BOTTOM, pady=14)


def main() -> None:
    root = tk.Tk()
    app = OppoCheckerGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
